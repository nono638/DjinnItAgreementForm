"""The trial run sheet: which reporter wrote which pages, take by take, kept as an Excel workbook.

When several reporters share a trial, each needs to know how many pages they wrote, for billing. The run
sheet has a row per take, with the columns of the reporters' own Google Sheet:

  Date | Weekday | Reporter | Day's Take | Running Take | Pages written | Starting Page No. |
  Ending Page No. | Witness Start | Witness End | Note

The weekday, the take counts and the page numbers are formulas ("Don't write here"), so a row typed in or
changed by hand is counted too. Filtering the Reporter column shows that reporter's pages (the total of the
rows shown is at the top), and the "By Reporter" sheet adds up each reporter's takes and pages.

There is one run sheet per case. A transcript's takes are added to the case's run sheet when there is one in
the run sheets folder (or next to the transcript), found by index number or, since one trial can have
several index numbers billed together, by case name; otherwise a new one is started. On a run sheet of this
app, the rows already there stay in their order (with what was typed in them) and each new take goes before
the first later one. A run sheet this app did not make (downloaded from Google Sheets, say) gets its rows
added at the end, or on a row typed in ahead for that day, with its own formulas copied down; it is backed up
first.
"""
from __future__ import annotations

import ast
import os
import re
import shutil
import threading
import time
from copy import copy
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from .dates import us_date
from .extract_regex import find_dates, norm_index, strip_line_numbers, tidy_name, title_page_count
from .fill import safe_filename, unique_path
from .models import CaseInfo
from .settings import Settings
from .takes import TITLE_ONLY, body_pages, find_takes, reporter_label, title_reporters

SHEET, BY_REPORTER, META = "Run Sheet", "By Reporter", "DjinnIt"
FORMAT = "DjinnIt run sheet 1"
HEAD_ROW, FIRST = 4, 5  # the column headings, the first take
LAST = 5000             # how far down the totals look
# key, heading, typed by hand ("Write here"), width
COLUMNS = [
    ("date", "Date", True, 11), ("weekday", "Weekday", False, 11), ("reporter", "Reporter", True, 13),
    ("day_take", "Day's Take", False, 9), ("run_take", "Running Take", False, 9),
    ("pages", "Pages written", True, 9), ("start", "Starting Page No.", False, 10),
    ("end", "Ending Page No.", False, 10), ("witness_start", "Witness Start", True, 24),
    ("witness_end", "Witness End", True, 24), ("note", "Note", True, 34),
]
COL = {k: i + 1 for i, (k, *_rest) in enumerate(COLUMNS)}
L = {k: chr(64 + i) for k, i in COL.items()}  # column letters
GREEN, ORANGE, RED, BAND = "93C47D", "F6B26B", "EA9999", "E0F7FA"
NO_RUNSHEET = "a run sheet needs a transcript PDF (its pages and the reporters' initials)"
NOT_FOUND = "reporter's initials not found"
HAND = [k for k, _, hand, _ in COLUMNS if hand]  # the columns typed by hand
_LOCK = threading.Lock()


@dataclass
class Row:
    """One take on the run sheet."""
    day: date | None = None
    reporter: str = ""
    pages: int | None = None
    start: int | None = None
    witness_start: str = ""
    witness_end: str = ""
    note: str = ""
    initials: str = ""  # as on the pages ("ds"); not a column of its own
    extra: dict = field(default_factory=dict)  # cells to the right of the run sheet's columns, kept with the row
    raw: dict = field(default_factory=dict)    # typed cells kept as they were (a formula, "TBD"), by column key
    origin: int | None = None                  # the row it was on (its formulas are moved with it)

    def end(self) -> int | None:
        """The take's last page number, or None without a first page or page count."""
        return self.start + self.pages - 1 if self.start is not None and self.pages else None

    def same_take(self, other: "Row") -> bool:
        """Already on the run sheet: the same day and pages that overlap, or (without page numbers) the
        same day, reporter and page count."""
        if self.day != other.day:
            return False
        if None not in (self.start, other.start, self.end(), other.end()):
            return self.start <= other.end() and other.start <= self.end()
        return (self.reporter.lower(), self.pages) == (other.reporter.lower(), other.pages)


@dataclass
class RunSheetOpts:
    """The takes to add, where to add them, and (after generate) what happened."""
    rows: list[Row] = field(default_factory=list)
    target: str | None = None  # None: as Settings.runsheet_existing says; "": a new run sheet; else that file
    path: Path | None = None   # the run sheet written
    created: bool = False      # a new one
    added: int = 0
    skipped: int = 0           # takes that were on it already (or given twice)
    added_pages: int = 0       # the pages of the takes added


def run_sheet_summary(opts: RunSheetOpts) -> str:
    """What generate did to the run sheet, for the 'Saved' message."""
    name = opts.path.name if opts.path else "the run sheet"
    takes = lambda n: f"{n} take{'s' if n != 1 else ''}"
    if opts.created:
        text = f"Started the run sheet {name} with {takes(opts.added)}."
    elif opts.added:
        text = f"Added {takes(opts.added)} to the run sheet {name}."
    else:
        return f"The run sheet {name} already has these takes: nothing was added."
    if opts.skipped:
        text += f" ({takes(opts.skipped)} {'was' if opts.skipped == 1 else 'were'} on it already.)"
    return text


@dataclass
class Found:
    """A run sheet on disk that may be this case's."""
    path: Path
    case_name: str = ""
    index_nos: list[str] = field(default_factory=list)
    ours: bool = True  # made by this app (False: someone else's, added to as it is laid out)
    why: str = ""  # "index" (the same index number) or "name" (the same case name)

    def reason(self) -> str:
        """Why this may be the case's run sheet, for the user ("same index number")."""
        if self.why == "index":
            return "same index number"
        return "same case name" + (f", index {', '.join(self.index_nos)}" if self.index_nos else "")


# ------------------------------------------------------------------ transcripts -> rows

def rows_from(ing, day: date | None, s: Settings) -> list[Row]:
    """The run sheet rows of one transcript PDF: a row per take. ing: the ingested PDF (its page marks,
    page count and text); day: the day the transcript is of. The reporters are named as reporter_label says."""
    if not ing.marks:  # no text to read the initials from (a scan): one row for the whole transcript
        return [Row(day, "", ing.page_count, ing.first_page_no, note=NOT_FOUND)]
    title_text, _ = strip_line_numbers(ing.text)
    names = title_reporters("\n".join(title_text.split("\f")[:title_page_count(title_text)]))
    takes = find_takes(ing.marks, title_page_count(title_text))
    p = s.profile
    rows = []
    for t in takes:
        who = reporter_label(t.initials, names, s.reporters, p.name, p.initials, s.title_case_names)
        note = TITLE_ONLY if t.title_only and len(takes) > 1 else "" if t.initials else NOT_FOUND
        rows.append(Row(day, who, t.pages, t.start, "; ".join(tidy_name(w, s.title_case_names) for w in t.witness_start),
                        "; ".join(tidy_name(w, s.title_case_names) for w in t.witness_end), note, t.initials))
    return rows


def name_rows(rows: list[Row], known: dict[str, str] | None = None) -> dict[str, str]:
    """Gives every take of a reporter the same name. One transcript's title page may name a reporter where
    another's doesn't, and the run sheet would say "DS" on one day and "Dana" on the next. known: names
    learned before (initials -> name). Changes the rows in place and returns what is known now."""
    def initials(r: Row) -> str:
        return r.initials or (r.reporter.lower() if re.fullmatch(r"[A-Z]{2,3}", r.reporter) else "")
    names = dict(known or {})
    for r in rows:
        if initials(r) and r.reporter and r.reporter != initials(r).upper():
            names.setdefault(initials(r), r.reporter)
    for r in rows:
        if initials(r) and r.reporter == initials(r).upper() and initials(r) in names:
            r.reporter = names[initials(r)]
    return names


def transcript_pages(ing) -> int:
    """The pages of an ingested transcript that are written and billed (not the word index)."""
    return body_pages(ing.marks, ing.page_count)


def to_date(v) -> date | None:
    """A cell's date: a date, a datetime or 'M/D/YYYY' text."""
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        found = find_dates(v)
        if found:
            m, d, y = (int(x) for x in found[0][2].split("/"))
            try:
                return date(y, m, d)
            except ValueError:
                return None
    return None


# ------------------------------------------------------------------ which run sheet is this case's

# words that don't tell one case from another: party roles, file-name words, places, months
_MONTHS = {date(2000, m, 1).strftime("%B").lower() for m in range(1, 13)} | \
          {date(2000, m, 1).strftime("%b").lower() for m in range(1, 13)} | {"sept"}
_NOISE = {"the", "of", "and", "et", "al", "ano", "inc", "llc", "llp", "pc", "corp", "co", "md", "as", "by", "an",
          "administrator", "administratrix", "executor", "executrix", "estate", "deceased", "individually",
          "guardian", "ad", "litem", "infant", "his", "her", "their", "parent", "natural", "on", "behalf",
          "page", "pages", "sheet", "run", "transcript", "transcripts", "copy", "trial", "minutes", "new", "york",
          "city", "county", "state", "people"} | _MONTHS


def index_numbers(text: str) -> list[str]:
    """Every index number in text, as 'number/year' ('712345-2021' -> '712345/2021')."""
    out = []
    for num, yr in re.findall(r"(?<!\d)(\d{3,7})\s*[-/]\s*(\d{4}|\d{2})(?!\d)", text or ""):
        n = norm_index(num, yr)
        if n and n not in out:
            out.append(n)
    return out


# words of a run sheet's file name, and of a company's name, that are never the party's own
_SHEET_WORDS = {"page", "pages", "sheet", "run", "transcript", "transcripts", "copy", "trial", "minutes"}
_ENTITY = {"inc", "llc", "llp", "pllc", "lp", "pc", "corp", "corporation", "co", "company", "ltd", "md", "do", "dds",
           "esq"}


def _sides(name: str) -> list[str]:
    """One word per side of the 'v.' that tells the case apart: the last name of the first party ('Jane Roe,
    as Administrator ...' -> 'roe'), or the last telling word of a company's name."""
    parts = re.split(r"\s(?:v|vs|versus|against)\.?\s", " " + (name or "").lower().replace("-against-", " v ") + " ",
                     maxsplit=1)
    keys = []
    for p in parts:
        first = re.split(r",|\sas\s|\sand\s|&|\bet\s+al\b", p, maxsplit=1)[0]  # the first party
        words = [w for w in re.findall(r"[a-z][a-z']+", first) if w not in _NOISE and w not in _ENTITY]
        # "The City of New York": nothing but common words, so its last word (not "Run Sheet" of a file name)
        raw = [w for w in re.findall(r"[a-z][a-z']+", first) if w not in _SHEET_WORDS and w not in _MONTHS]
        keys.append(words[-1] if words else raw[-1] if raw else "")
    return keys


def same_name(a: str, b: str) -> bool:
    """'Jane Roe, as Administrator ... v. Sam Poe, M.D., ...' and 'ROE v Poe et al.': the first party on each
    side has the same last name."""
    x, y = _sides(a), _sides(b)
    return len(x) == len(y) == 2 and all(x) and x == y


def runsheets_folder(s: Settings) -> Path:
    """Where run sheets are kept: Settings.runsheet_dir, else Documents/DjinnIt Run Sheets."""
    return Path(s.runsheet_dir) if s.runsheet_dir else Path.home() / "Documents" / "DjinnIt Run Sheets"


def read_info(path: Path) -> Found | None:
    """What a workbook says about its case, or None when it isn't a run sheet (or can't be read)."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return None
    try:
        if META in wb.sheetnames:
            info = {}
            for key, value, *_ in wb[META].iter_rows(values_only=True):
                if key:
                    info.setdefault(str(key), []).append("" if value is None else str(value))
            if info.get("format", [""])[0].startswith("DjinnIt run sheet"):
                return Found(path, (info.get("case") or [""])[0], [i for i in info.get("index", []) if i])
        if {"Setup", "Rates", "Job", "Invoice"} <= set(wb.sheetnames):
            return None  # the invoice spreadsheet: its Page Log looks like a run sheet but isn't one
        for ws in wb.worksheets:  # someone else's run sheet: the columns are there, the case is in the name
            top = [[c for c in row if c is not None] for row in ws.iter_rows(max_row=12, values_only=True)]
            if _header_row(top) is not None:
                # only text: a date cell ("2026-06-02") is not an index number
                text = " ".join([path.stem] + [c for row in top[:3] for c in row if isinstance(c, str)])
                return Found(path, path.stem, index_numbers(text), ours=False)
    except Exception:
        return None
    finally:
        wb.close()
    return None


def _header_row(rows: list[list]) -> int | None:
    """The row (from 0) of a run sheet's column headings: one with Date, Reporter and Pages in it."""
    for i, row in enumerate(rows):
        cells = " | ".join(str(c).lower() for c in row if c is not None)
        if "reporter" in cells and "page" in cells and "date" in cells:
            return i
    return None


def matches(found: Found, case_name: str, index_no: str) -> str:
    """'index', 'name' or '' - how found is this case's run sheet."""
    if set(index_numbers(index_no)) & set(found.index_nos):
        return "index"
    return "name" if case_name and same_name(case_name, found.case_name) else ""


def find_sheets(case: CaseInfo, folders: list[Path]) -> list[Found]:
    """The run sheets in these folders that may be this case's: same index number first, then this app's own
    before others, then the newest. Excel's lock files, half-saved files and backups are skipped."""
    out, seen = [], set()
    for folder in folders:
        try:
            files = sorted(Path(folder).glob("*.xlsx"))
        except OSError:
            continue
        for f in files:
            key = str(f.resolve()).lower()
            if f.name.startswith("~$") or f.name.endswith(".saving.xlsx") or "(before DjinnIt)" in f.name \
                    or key in seen:
                continue
            seen.add(key)
            info = read_info(f)
            if info:
                info.why = matches(info, case.get("case_name"), case.get("index_no"))
                if info.why:
                    out.append(info)
    return sorted(out, key=lambda x: (x.why != "index", not x.ours, -x.path.stat().st_mtime))


def choose(case: CaseInfo, s: Settings, folders: list[Path], target: str | None) -> Path | None:
    """The run sheet to add to (None: start a new one). target: what the user chose (see RunSheetOpts);
    without a choice, Settings.runsheet_existing decides - "ask" (nobody to ask here) adds to a run sheet
    with the same index number only."""
    if target is not None:
        return Path(target) if target else None
    if s.runsheet_existing == "new":
        return None
    found = find_sheets(case, folders)
    if s.runsheet_existing != "add":
        found = [f for f in found if f.why == "index"]
    return found[0].path if found else None


# ------------------------------------------------------------------ writing

def new_path(case: CaseInfo, rows: list[Row], s: Settings) -> Path:
    """A free file name for a new run sheet in the run sheets folder, from Settings.runsheet_filename_pattern
    and the first day of the takes ("June 2026 712345-2021 Jane Roe v. Sam Poe - Run Sheet.xlsx")."""
    days = [r.day for r in rows if r.day]
    first = min(days) if days else date.today()
    from .fill import short_caption
    try:
        name = s.runsheet_filename_pattern.format(
            month=first.strftime("%B"), year=first.year, case=short_caption(case.get("case_name")) or "Case",
            index=case.get("index_no").replace("/", "-") or "no index", date=us_date(first).replace("/", "-"))
    except (KeyError, IndexError, ValueError, AttributeError, TypeError):  # a bad custom pattern
        name = f"Run Sheet - {case.get('index_no').replace('/', '-')}"
    name = re.sub(r"\s+", " ", name).strip(" -")
    return unique_path(runsheets_folder(s) / (safe_filename(name or "Run Sheet") + ".xlsx"))


def add_takes(case: CaseInfo, opts: RunSheetOpts, s: Settings, folders: list[Path]) -> Path:
    """Adds opts.rows to this case's run sheet (see choose) or a new one, and returns its path. folders: where
    else to look besides the run sheets folder (the transcript's own folder). Sets opts.path, created, added,
    skipped and added_pages. ValueError when there are no rows or the chosen file isn't a run sheet."""
    if not opts.rows:
        raise ValueError(NO_RUNSHEET)
    with _LOCK:  # one change at a time: the window and a batch may add to the same run sheet
        path = choose(case, s, [runsheets_folder(s), *folders], opts.target)
        if path and not path.exists():  # chosen, then moved or deleted: a new one in its place
            path = None
        info = read_info(path) if path else None
        if path and not info:
            raise ValueError(f"{path.name} is not a run sheet (no Date, Reporter and Pages columns)")
        if info and not info.ours:
            _append_theirs(path, opts)
        else:
            path = path or new_path(case, opts.rows, s)
            opts.created = info is None
            _write_ours(path, case, opts, exists=info is not None)
        opts.path = path
        return path


def _save(wb, path: Path) -> None:
    """Saves in one step: a run sheet open in Excel stays as it was (PermissionError). Dropbox or OneDrive may
    hold the file for a moment while they sync it, so a locked file is tried again a few times first."""
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.stem + ".saving.xlsx")
    try:
        wb.save(tmp)
        for attempt in range(3):
            try:
                os.replace(tmp, path)
                break
            except PermissionError:
                if attempt == 2:
                    raise
                time.sleep(0.3)
    except PermissionError as e:
        raise PermissionError(f"{path.name} is open in another program - close it in Excel and try again") from e
    finally:
        tmp.unlink(missing_ok=True)


def _number(v) -> int | None:
    """A cell's whole number: 12, 12.0, "12" or a sum typed in ("=329-320+1"); None otherwise."""
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return int(v)
    if isinstance(v, str) and v.strip().isdigit():
        return int(v)
    if isinstance(v, str) and re.fullmatch(r"=[\d\s+\-*()]+", v):
        try:
            return int(_sum(ast.parse(v[1:].strip(), mode="eval").body))
        except (SyntaxError, ValueError, TypeError, RecursionError):
            return None
    return None


def _sum(node) -> int:
    """The value of a sum of whole numbers (+, - and * only)."""
    ops = {ast.Add: lambda a, b: a + b, ast.Sub: lambda a, b: a - b, ast.Mult: lambda a, b: a * b}
    if isinstance(node, ast.Constant) and type(node.value) is int:
        return node.value
    if isinstance(node, ast.BinOp) and type(node.op) in ops:
        return ops[type(node.op)](_sum(node.left), _sum(node.right))
    if isinstance(node, ast.UnaryOp) and isinstance(node.op, ast.USub):
        return -_sum(node.operand)
    raise ValueError("not a sum")


def _unique(rows: list[Row]) -> list[Row]:
    """The rows less any that repeat an earlier one (the same transcript dropped twice)."""
    out: list[Row] = []
    for r in rows:
        if not any(r.same_take(o) for o in out):
            out.append(r)
    return out


def _read_ours(ws) -> list[Row]:
    """The takes on a run sheet of this app, top to bottom; rows with nothing typed in them are skipped."""
    rows, prev_end = [], None
    for r in range(FIRST, ws.max_row + 1):
        v = {k: ws.cell(r, c).value for k, c in COL.items()}
        extra = {c: ws.cell(r, c).value for c in range(len(COLUMNS) + 1, ws.max_column + 1)
                 if ws.cell(r, c).value is not None}
        if all(v[k] in (None, "") for k in HAND) and not extra:
            continue
        pages = _number(v["pages"])
        start = v["start"]
        if not isinstance(start, (int, float)) or isinstance(start, bool):  # a formula: the page after the row above
            start = prev_end + 1 if prev_end is not None else None
        day = to_date(v["date"])
        # what was typed is written back as it was: a formula, or a date or page count that isn't one ("TBD")
        raw = {k: v[k] for k in HAND if _is_formula(v[k])}
        if v["date"] not in (None, "") and day is None:
            raw["date"] = v["date"]
        if v["pages"] not in (None, "") and pages is None:
            raw["pages"] = v["pages"]
        row = Row(day, str(v["reporter"] or ""), pages, int(start) if start is not None else None,
                  str(v["witness_start"] or ""), str(v["witness_end"] or ""), str(v["note"] or ""), extra=extra,
                  raw=raw, origin=r)
        prev_end = row.end() if row.end() is not None else prev_end
        rows.append(row)
    return rows


def _in_order(old: list[Row], new: list[Row]) -> list[Row]:
    """The rows of the run sheet: the old ones where they were (in the order the user keeps them), each new one
    before the first old row that comes later by date, then first page. Old rows without a date or first page
    are passed over."""
    key = lambda r: (r.day or date.max, r.start if r.start is not None else float("inf"))
    rows = list(old)
    olds = {id(r) for r in old}
    for n in sorted(new, key=key):
        at = next((i for i, o in enumerate(rows) if id(o) in olds and o.day and o.start is not None
                   and key(o) > key(n)), len(rows))
        rows.insert(at, n)
    return rows


def _write_ours(path: Path, case: CaseInfo, opts: RunSheetOpts, exists: bool) -> None:
    """Writes a run sheet of this app: a new one, or (exists) the one at path rewritten with its old rows and
    the new takes in order. Sets opts.added, skipped and added_pages."""
    from openpyxl import Workbook, load_workbook
    if exists:
        wb = load_workbook(path)
        ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]
        old = _read_ours(ws)
        meta = {}
        if META in wb.sheetnames:
            for key, value, *_ in wb[META].iter_rows(values_only=True):
                if key:
                    meta.setdefault(str(key), []).append("" if value is None else str(value))
        case_name = (meta.get("case") or [case.get("case_name")])[0] or case.get("case_name")
        indexes = [i for i in meta.get("index", []) if i]
        known = dict(x.split("=", 1) for x in meta.get("reporter", []) if "=" in x)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = SHEET
        old, case_name, indexes, known = [], case.get("case_name"), [], {}
    new = [r for r in _unique(opts.rows) if not any(r.same_take(o) for o in old)]
    opts.added, opts.skipped = len(new), len(opts.rows) - len(new)
    opts.added_pages = sum(r.pages or 0 for r in new)
    if exists and not new:
        return  # nothing to add: the run sheet stays as it is (it may be open in Excel)
    for i in index_numbers(case.get("index_no")):
        if i not in indexes:
            indexes.append(i)
    rows = _in_order(old, new)
    known = name_rows(rows, known)
    _fill_sheet(ws, rows, case_name, indexes, case)
    _by_reporter(wb, rows, ws.title)
    _meta(wb, case_name, indexes, known)
    _save(wb, path)


def _moved(value, origin: int | None, row: int, col: int):
    """A formula written on another row than it was on: its references move with it (=A5 on row 5 is =A7 on
    row 7), as when Excel moves a row."""
    if not _is_formula(value) or origin in (None, row):
        return value
    from openpyxl.formula.translate import Translator
    from openpyxl.utils import get_column_letter
    letter = get_column_letter(col)
    try:
        return Translator(value, origin=f"{letter}{origin}").translate_formula(f"{letter}{row}")
    except Exception:  # a formula openpyxl can't read: as it was
        return value


def _fill_sheet(ws, rows: list[Row], case_name: str, indexes: list[str], case: CaseInfo) -> None:
    """Writes the whole Run Sheet tab again: the case at the top, the totals of the rows shown, the headings,
    and a row per take with its formulas."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.properties import PageSetupProperties
    for rng in list(ws.merged_cells.ranges):  # merged cells can't be written: unmerged where the sheet is rewritten
        if rng.min_col <= len(COLUMNS) or rng.max_row >= FIRST:
            ws.unmerge_cells(str(rng))
    for r in range(1, max(ws.max_row, FIRST) + 1):  # start over: the rows are written again in order
        # (the user's own cells to the right of the headings stay; those of the takes move with their rows)
        for c in range(1, (len(COLUMNS) if r < FIRST else max(ws.max_column, len(COLUMNS))) + 1):
            cell = ws.cell(r, c)
            cell.value = None
            cell.style = "Normal"
    fill = lambda color: PatternFill("solid", start_color=color, end_color=color)
    thin = Side(style="thin", color="BFBFBF")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    from .fill import short_caption
    ws["A1"] = short_caption(case_name, 120) or "Run sheet"
    ws["A1"].font = Font(size=14, bold=True)
    details = [f"Index No. {', '.join(indexes)}" if indexes else ""]
    if case.get("judge"):
        details.append(f"Justice {case.get('judge')}")
    if case.get("part"):
        details.append(f"Part {case.get('part')}")
    ws["A2"] = "   ·   ".join(d for d in details if d)
    ws["A2"].font = Font(color="595959")
    # the totals of the rows shown: filter the Reporter column to see one reporter's pages
    ws[f"{L['witness_start']}2"] = "Pages shown:"
    ws[f"{L['witness_start']}2"].alignment = Alignment(horizontal="right")
    ws[f"{L['witness_end']}2"] = f"=SUBTOTAL(109,{L['pages']}{FIRST}:{L['pages']}{LAST})"
    ws[f"{L['witness_end']}2"].font = Font(bold=True)
    ws[f"{L['witness_end']}2"].alignment = Alignment(horizontal="left")
    ws[f"{L['note']}2"] = f'="Rows shown: "&SUBTOTAL(103,{L["reporter"]}{FIRST}:{L["reporter"]}{LAST})'
    ws[f"{L['note']}2"].font = Font(color="595959")
    for key, heading, hand, width in COLUMNS:
        c = COL[key]
        mark = ws.cell(3, c, "Write here" if hand else "Don't write here")
        mark.fill, mark.font = fill(GREEN if hand else RED), Font(size=6, bold=True, color="FFFFFF" if hand else "000000")
        mark.alignment = Alignment(horizontal="center")
        head = ws.cell(HEAD_ROW, c, heading)
        head.fill, head.font, head.border = fill(GREEN if hand else ORANGE), Font(bold=True), box
        head.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[L[key]].width = width
    ws.row_dimensions[HEAD_ROW].height = 30

    prev_end, prev_end_row, prev_day = None, None, None
    for i, row in enumerate(rows):
        r = FIRST + i
        a, d, f, g = (f"{L[k]}{r}" for k in ("date", "day_take", "pages", "start"))
        above = r - 1
        values = {
            "date": row.day, "reporter": row.reporter, "pages": row.pages,
            "witness_start": row.witness_start or None, "witness_end": row.witness_end or None, "note": row.note or None,
            "weekday": f'=IF({a}="","",TEXT({a},"dddd"))',
            "day_take": f'=IF({a}="","",IF({L["note"]}{r}="{TITLE_ONLY}",0.5,'
                        f'COUNTIFS({L["date"]}${FIRST}:{a},{a},{L["note"]}${FIRST}:{L["note"]}{r},"<>{TITLE_ONLY}")))',
            "run_take": f'=IF({a}="","",N({L["run_take"]}{above})+IF({a}={L["date"]}{above},{d}-N({L["day_take"]}{above}),{d}))',
            "end": f'=IF(OR({f}="",{g}=""),"",{g}+{f}-1)',
        }
        if row.start is not None and prev_end is not None and row.start == prev_end + 1:
            values["start"] = f"={L['end']}{prev_end_row}+1"  # the row the last page came from (not a blank one)
        else:
            values["start"] = row.start
        for key, value in row.raw.items():
            values[key] = _moved(value, row.origin, r, COL[key])
        first_of_day = row.day != prev_day
        for key, *_ in COLUMNS:
            cell = ws.cell(r, COL[key], values[key])
            cell.border = box
            if i % 2:
                cell.fill = fill(BAND)
            if key in ("date", "start") and first_of_day:
                cell.font = Font(bold=True)
            if key in ("weekday", "day_take", "run_take", "pages", "start", "end"):
                cell.alignment = Alignment(horizontal="center")
        ws.cell(r, COL["date"]).number_format = "m/d/yyyy"
        for c, value in row.extra.items():
            ws.cell(r, c, _moved(value, row.origin, r, c))
        if row.end() is not None:
            prev_end, prev_end_row = row.end(), r
        prev_day = row.day
    ws.freeze_panes = f"A{FIRST}"
    ws.auto_filter.ref = f"A{HEAD_ROW}:{L['note']}{max(FIRST, FIRST + len(rows) - 1)}"
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.print_title_rows = f"{HEAD_ROW}:{HEAD_ROW}"


def _by_reporter(wb, rows: list[Row], sheet: str = SHEET) -> None:
    """The 'By Reporter' sheet: each reporter's takes and pages (formulas over the run sheet, the tab named
    sheet)."""
    from openpyxl.styles import Font
    if BY_REPORTER in wb.sheetnames:
        del wb[BY_REPORTER]
    ws = wb.create_sheet(BY_REPORTER, 1)
    for c, (heading, width) in enumerate((("Reporter", 18), ("Takes", 9), ("Pages", 9)), 1):
        ws.cell(1, c, heading).font = Font(bold=True)
        ws.column_dimensions[chr(64 + c)].width = width
    tab = "'" + sheet.replace("'", "''") + "'"
    ref = lambda k: f"{tab}!${L[k]}${FIRST}:${L[k]}${LAST}"
    names: dict[str, str] = {}  # COUNTIF and SUMIF don't tell "Dana" from "DANA": one row for both
    for r in rows:
        if r.reporter and not _is_formula(r.reporter):
            names.setdefault(r.reporter.casefold(), r.reporter)
    for i, name in enumerate(names.values(), 2):
        ws.cell(i, 1, name)
        ws.cell(i, 2, f'=COUNTIFS({ref("reporter")},A{i},{ref("note")},"<>{TITLE_ONLY}")')
        ws.cell(i, 3, f'=SUMIF({ref("reporter")},A{i},{ref("pages")})')
    total = len(names) + 2
    ws.cell(total, 1, "Total").font = Font(bold=True)
    ws.cell(total, 2, f"=SUM(B2:B{total - 1})").font = Font(bold=True)
    ws.cell(total, 3, f"=SUM(C2:C{total - 1})").font = Font(bold=True)


def _meta(wb, case_name: str, indexes: list[str], reporters: dict[str, str]) -> None:
    """The hidden sheet that says what the run sheet is about (so it can be found again)."""
    if META in wb.sheetnames:
        del wb[META]
    ws = wb.create_sheet(META)
    ws.append(["format", FORMAT])
    ws.append(["case", case_name])
    for i in indexes or [""]:
        ws.append(["index", i])
    for initials, name in sorted(reporters.items()):
        ws.append(["reporter", f"{initials}={name}"])
    ws.append(["updated", datetime.now().isoformat(timespec="seconds")])
    ws.sheet_state = "hidden"


# ------------------------------------------------------------------ someone else's run sheet

# headings (spaces and line breaks made single spaces, lower case); the first that fits a heading is its column
_THEIR_KEYS = [("date", r"\bdates?\b"), ("weekday", r"weekday|day of"), ("reporter", r"reporter"),
               ("day_take", r"day.?s take"), ("run_take", r"running take"), ("start", r"start(ing)? page"),
               ("end", r"end(ing)? page"), ("pages", r"\bpages?\b"), ("witness_start", r"witness start"),
               ("witness_end", r"witness end"), ("note", r"^notes?$")]


def _their_columns(ws) -> tuple[int, dict[str, int]]:
    """The row of the column headings (from 1) and the column of each heading found, by key (see COLUMNS)."""
    rows = [[c.value for c in row] for row in ws.iter_rows(min_row=1, max_row=12)]
    h = _header_row(rows)
    if h is None:
        raise ValueError(f"{ws.title}: no Date, Reporter and Pages columns")
    cols = {}
    for c, value in enumerate(rows[h], 1):
        text = " ".join(str(value or "").split()).lower()
        for key, pattern in _THEIR_KEYS:
            if key not in cols and text and re.search(pattern, text):
                cols[key] = c
                break
    missing = [name for key, name in (("date", "Date"), ("reporter", "Reporter"), ("pages", "Pages")) if key not in cols]
    if missing:
        raise ValueError(f"{ws.title}: no {' or '.join(missing)} column found")
    return h + 1, cols


def _is_formula(v) -> bool:
    return isinstance(v, str) and v.startswith("=")


def _append_theirs(path: Path, opts: RunSheetOpts) -> None:
    """Adds the new takes to a run sheet this app did not make, keeping its layout: each on a row typed in
    ahead for its day, else on the first empty row after the takes, with the look and formulas of the row
    above. The file is first copied to "<name> (before DjinnIt).xlsx". Sets opts.added, skipped and
    added_pages."""
    from openpyxl import load_workbook
    from openpyxl.formula.translate import Translator
    wb = load_workbook(path)
    cached = load_workbook(path, data_only=True)  # the formulas' last values, as the spreadsheet saved them
    ws = next(w for w in wb.worksheets if _header_row([[c.value for c in r] for r in w.iter_rows(max_row=12)])
              is not None)
    vals = cached[ws.title]
    head, cols = _their_columns(ws)
    get = lambda r, k: ws.cell(r, cols[k]).value if k in cols else None
    num = lambda r, k: (lambda v: v if isinstance(v, (int, float)) and not isinstance(v, bool) else None)(
        vals.cell(r, cols[k]).value if k in cols else None)

    def day_at(r: int) -> date | None:
        v = get(r, "date")
        return to_date(vals.cell(r, cols["date"]).value) if _is_formula(v) else to_date(v)

    def pages_at(r: int) -> int | None:
        p = get(r, "pages")
        return int(p) if isinstance(p, (int, float)) and not isinstance(p, bool) and p > 0 else None

    # what is on it: the rows with a reporter or pages (the page numbers worked out where not saved)
    old, last, prev_end = [], head, None
    ends: dict[int, int] = {}  # the last page of each row that has one
    for r in range(head + 1, ws.max_row + 1):
        pages = pages_at(r)
        if not get(r, "reporter") and pages is None:  # an empty row, or a day typed in ahead
            continue
        start = get(r, "start")
        start = start if isinstance(start, (int, float)) and not _is_formula(start) else num(r, "start")
        if start is None and prev_end is not None:
            start = prev_end + 1
        row = Row(day_at(r), str(get(r, "reporter") or ""), pages, int(start) if start else None)
        end = num(r, "end")
        end = row.end() if row.end() is not None else int(end) if end is not None else None
        if end is not None:
            ends[r] = prev_end = end
        old.append(row)
        last = r
    new = [x for x in _unique(opts.rows) if not any(x.same_take(o) for o in old)]
    opts.added, opts.skipped = len(new), len(opts.rows) - len(new)
    opts.added_pages = sum(r.pages or 0 for r in new)
    if not new:
        return
    backup = path.with_name(f"{path.stem} (before DjinnIt){path.suffix}")
    if not backup.exists():
        shutil.copy2(path, backup)
    typed = [cols.get(k) for k in ("date", "reporter", "pages", "witness_start", "witness_end", "note")]
    used: set[int] = set()

    def free(r: int) -> bool:  # no take on it, nor one of this run
        return r not in used and not get(r, "reporter") and pages_at(r) is None

    for row in new:
        # the row typed in ahead for this take's day, else the first empty row after the takes
        r = next((x for x in range(head + 1, ws.max_row + 1) if row.day and free(x) and day_at(x) == row.day), None)
        if r is None:
            r = next(x for x in range(last + 1, ws.max_row + 2 + len(new)) if free(x) and day_at(x) is None)
        for c in range(1, ws.max_column + 1):  # the look of the row above, its formulas copied down
            cell, above = ws.cell(r, c), ws.cell(r - 1, c)
            if above.has_style and not cell.has_style:
                cell._style = copy(above._style)
            if cell.value in (None, "") and _is_formula(above.value) and c not in typed:
                cell.value = Translator(above.value, origin=above.coordinate).translate_formula(cell.coordinate)
        for key, value in (("date", row.day), ("reporter", row.reporter), ("pages", row.pages),
                           ("witness_start", row.witness_start), ("witness_end", row.witness_end),
                           ("note", row.note)):
            # a date formula that already gives this day is kept
            if key in cols and value not in (None, "") and not (key == "date" and _is_formula(get(r, key))
                                                              and day_at(r) == row.day):
                ws.cell(r, cols[key]).value = value
        if "start" in cols:
            cell = ws.cell(r, cols["start"])
            follows = ends.get(r - 1) is not None and row.start == ends[r - 1] + 1
            if row.start is not None and not (follows and _is_formula(cell.value)):
                cell.value = row.start  # a new numbering, no formula to work it out, or nothing above to go on
            elif row.start is None and _is_formula(cell.value) and ends.get(r - 1) is None:
                cell.value = None       # it would point at a row without pages
        if "end" in cols and not _is_formula(ws.cell(r, cols["end"]).value) and row.end() is not None:
            ws.cell(r, cols["end"]).value = row.end()
        if "weekday" in cols and not _is_formula(ws.cell(r, cols["weekday"]).value) and row.day:
            ws.cell(r, cols["weekday"]).value = row.day.strftime("%A")
        if row.end() is not None:
            ends[r] = row.end()
        used.add(r)
        last = max(last, r)
    _save(wb, path)
