"""Records of everything the app made, and the invoice ledger.

Kept in a small SQLite database (%APPDATA%\\DjinnItAgreementForm\\records.db, the program's own copy) and
mirrored as invoices.csv / activity.csv in the records folder (Settings.records_folder) after every
change, so the data can always be opened in Excel. Unlike the log file, these hold case details:
they are the user's own business records and never leave the computer.

  activity  one row per file made (minute agreement, MOFR, invoice) or run sheet added to
  invoices  one row per invoice, with the amount of each speed offered and whether it was paid

An invoice's "billed" amount is what it was paid at once paid, else the price of the first speed it
offers: the job's own speed, or on a choice invoice the first (slowest, cheapest) of the speeds it offers
(the attorney picks; that is what they owe at the least). Void ones count 0. The amounts are the ones the
invoice was made with, until they are changed here (Ledger.set_amounts: an amount corrected in the PDF).

A record deleted goes to the trash (its `deleted` time is set): it is left out of every list, total and copy,
and can be restored for 30 days (TRASH_DAYS); after that it is deleted for good. The files themselves are
never touched. An invoice's number stays taken even then (the used_numbers table), so it is never given to
another invoice.

Databases made by older versions get the newer columns when opened (_migrate); their old rows leave them blank.
"""
from __future__ import annotations

import csv
import html
import json
import sqlite3
import threading
from contextlib import closing
from dataclasses import asdict, dataclass, field, fields
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path

from .dates import us_date
from .invoice_calc import fmt, money
from .log import error as log_error
from .rates import parse_amount

# One writer at a time: the CSV copies are rewritten whole, and an invoice's number is taken and its row
# added in one step (the batch runs on another thread than the window).
_MIRROR_LOCK = threading.Lock()
NUMBER_LOCK = threading.RLock()
# Opening a database makes its tables and adds missing columns: two threads doing that at once (the batch and the
# window) would both try to add the same column
_SCHEMA_LOCK = threading.Lock()
KINDS = {"agreement": "Minute agreement", "mofr": "MOFR", "invoice": "Invoice", "runsheet": "Run sheet"}
TRASH_DAYS = 30  # a deleted record can be restored for this long
SCHEMA_VERSION = 2
# Columns added after the first version: (table, column, type), added to an older database when it is opened
# (a new one gets them from _SCHEMA)
_ADDED = [
    ("invoices", "court", "TEXT"), ("invoices", "part", "TEXT"), ("invoices", "transcript_pages", "INTEGER"),
    ("invoices", "my_pages", "INTEGER"), ("invoices", "reporters", "TEXT"), ("invoices", "excerpt", "TEXT"),
    ("invoices", "email_copy", "TEXT"), ("invoices", "idx", "TEXT"), ("invoices", "deleted", "TEXT"),
    ("activity", "transcript_pages", "INTEGER"), ("activity", "my_pages", "INTEGER"), ("activity", "deleted", "TEXT"),
]

_SCHEMA = """
CREATE TABLE IF NOT EXISTS activity (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    ts TEXT NOT NULL,             -- 2026-09-30T14:05:00
    kind TEXT NOT NULL,           -- agreement / mofr / invoice / runsheet
    case_name TEXT, index_no TEXT, dates TEXT, judge TEXT, part TEXT,
    attorney TEXT, firm TEXT, pages INTEGER, file_path TEXT, invoice_no TEXT,
    transcript_pages INTEGER, my_pages INTEGER, deleted TEXT
);
CREATE TABLE IF NOT EXISTS invoices (
    invoice_no TEXT PRIMARY KEY,
    year INTEGER, seq INTEGER,
    created TEXT NOT NULL,        -- 2026-09-30
    case_name TEXT, index_no TEXT, dates TEXT, judge TEXT,
    bill_to TEXT, firm TEXT, email TEXT,
    pages INTEGER, parties INTEGER,
    amounts TEXT NOT NULL,        -- JSON: speed -> amount each party pays, e.g. {"Regular": "63.00"}
    billed_speed TEXT,
    status TEXT NOT NULL DEFAULT 'open',
    paid_speed TEXT, amount_paid TEXT, paid_date TEXT,
    file_path TEXT, notes TEXT,
    court TEXT, part TEXT, transcript_pages INTEGER, my_pages INTEGER, reporters TEXT, excerpt TEXT,
    email_copy TEXT, idx TEXT, deleted TEXT
);
CREATE TABLE IF NOT EXISTS used_numbers (     -- every invoice number given, even of invoices deleted for good
    invoice_no TEXT PRIMARY KEY, year INTEGER, seq INTEGER
);
"""


def default_db() -> Path:
    """records.db in the app's settings folder."""
    from .settings import settings_dir
    return settings_dir() / "records.db"


@dataclass
class Invoice:
    """One invoice in the ledger. amounts: speed -> what each party pays ({"Regular": "63.00"});
    status: "open", "paid" or "void"."""
    invoice_no: str
    created: str
    case_name: str = ""
    index_no: str = ""
    dates: str = ""
    judge: str = ""
    bill_to: str = ""
    firm: str = ""
    email: str = ""
    pages: int = 0
    parties: int = 1
    amounts: dict[str, str] = field(default_factory=dict)
    billed_speed: str = ""
    status: str = "open"
    paid_speed: str = ""
    amount_paid: str = ""
    paid_date: str = ""
    file_path: str = ""
    notes: str = ""
    court: str = ""
    part: str = ""
    transcript_pages: int = 0      # every page of the transcript(s) billed, whoever wrote them
    my_pages: int = 0              # the user's own pages of them (fewer when several reporters wrote them)
    reporters: str = ""            # who wrote how many pages: "PR 65, DS 85"
    excerpt: str = ""              # the pages this firm ordered, when not all of them: "6/3/2026 pp. 20-40"
    email_copy: str = ""           # "Yes" / "No": an e-mailed copy was charged ("" = not known, older records)
    index: str = ""                # "Yes" / "No": an index was charged (the column is "idx")
    deleted: str = ""              # when it went to the trash (ISO time); "" = not deleted

    @property
    def client(self) -> str:
        """Who is billed, for grouping: the firm, else the attorney."""
        return self.firm or self.bill_to or "(no name)"

    @property
    def billed(self) -> Decimal:
        """What the invoice counts for (see the module docstring)."""
        if self.status == "void":
            return Decimal("0.00")
        if self.status == "paid":
            return money(self.amount_paid)
        return money(self.amounts.get(self.billed_speed) or next(iter(self.amounts.values()), "0"))

    @property
    def paid(self) -> Decimal:
        """What was paid: the amount paid once the invoice is paid, else 0."""
        return money(self.amount_paid) if self.status == "paid" else Decimal("0.00")

    @property
    def outstanding(self) -> Decimal:
        """What is still owed: the billed amount while the invoice is open, else 0."""
        return self.billed if self.status == "open" else Decimal("0.00")

    def offered_text(self) -> str:
        """'Regular $63.00; Expedite $95.00' (the speeds offered, for the tables)."""
        return "; ".join(f"{k} {fmt(money(v))}" for k, v in self.amounts.items())


@dataclass
class Activity:
    """One file made, or a run sheet added to (pages: the pages of the takes added). kind: a key of KINDS."""
    kind: str
    ts: str = ""                  # 2026-09-30T14:05:00; blank = now
    id: int = 0
    case_name: str = ""
    index_no: str = ""
    dates: str = ""
    judge: str = ""
    part: str = ""
    attorney: str = ""
    firm: str = ""
    pages: int = 0
    file_path: str = ""
    invoice_no: str = ""
    transcript_pages: int = 0     # every page of the transcript(s), whoever wrote them
    my_pages: int = 0             # the user's own pages of them
    deleted: str = ""             # when it went to the trash (ISO time); "" = not deleted


def gone_for_good(deleted: str) -> str:
    """The day a record deleted at `deleted` (ISO time) leaves the trash for good, as M/D/YYYY."""
    try:
        return us_date((datetime.fromisoformat(deleted) + timedelta(days=TRASH_DAYS)).date().isoformat())
    except ValueError:
        return ""


@dataclass
class Summary:
    """The count and money totals of a group of invoices (void ones are not counted)."""
    count: int = 0
    billed: Decimal = Decimal("0.00")
    paid: Decimal = Decimal("0.00")
    outstanding: Decimal = Decimal("0.00")

    def add(self, inv: Invoice) -> None:
        if inv.status != "void":
            self.count += 1
        self.billed += inv.billed
        self.paid += inv.paid
        self.outstanding += inv.outstanding


def summarize(invoices: list[Invoice]) -> tuple[Summary, dict[str, Summary], dict[str, Summary]]:
    """(totals, per client, per month 'YYYY-MM'); void invoices are left out of the counts and sums."""
    total, by_client, by_month = Summary(), {}, {}
    for inv in invoices:
        total.add(inv)
        if inv.status == "void":
            continue
        by_client.setdefault(inv.client, Summary()).add(inv)
        by_month.setdefault(inv.created[:7], Summary()).add(inv)
    return total, dict(sorted(by_client.items(), key=lambda kv: -kv[1].billed)), dict(sorted(by_month.items()))


# The exported tables (CSV, Excel): (header, value) per column
INVOICE_TABLE = [
    ("Invoice No.", lambda i: i.invoice_no), ("Date", lambda i: us_date(i.created)), ("Case", lambda i: i.case_name),
    ("Index No.", lambda i: i.index_no), ("Court", lambda i: i.court), ("Part", lambda i: i.part),
    ("Date(s) of proceeding", lambda i: i.dates), ("Judge", lambda i: i.judge),
    ("Bill to", lambda i: i.bill_to), ("Firm", lambda i: i.firm), ("E-mail", lambda i: i.email),
    ("Pages", lambda i: i.pages), ("My pages", lambda i: i.my_pages or ""),
    ("Transcript pages", lambda i: i.transcript_pages or ""), ("Reporters", lambda i: i.reporters),
    ("Excerpt", lambda i: i.excerpt), ("Parties", lambda i: i.parties),
    ("Speeds offered", lambda i: ", ".join(i.amounts)), ("Offered (per party)", Invoice.offered_text),
    ("E-mailed copy", lambda i: i.email_copy), ("Index", lambda i: i.index),
    ("Billed", lambda i: float(i.billed)), ("Status", lambda i: i.status.title()),
    ("Paid speed", lambda i: i.paid_speed), ("Amount paid", lambda i: float(i.paid) if i.status == "paid" else ""),
    ("Date paid", lambda i: us_date(i.paid_date)), ("File", lambda i: i.file_path), ("Notes", lambda i: i.notes),
]
ACTIVITY_TABLE = [
    ("Date", lambda a: us_date(a.ts)), ("Time", lambda a: a.ts[11:16]), ("Made", lambda a: KINDS.get(a.kind, a.kind)),
    ("Case", lambda a: a.case_name), ("Index No.", lambda a: a.index_no), ("Date(s) of proceeding", lambda a: a.dates),
    ("Judge", lambda a: a.judge), ("Part", lambda a: a.part), ("Attorney", lambda a: a.attorney),
    ("Firm", lambda a: a.firm), ("Pages", lambda a: a.pages or ""), ("My pages", lambda a: a.my_pages or ""),
    ("Transcript pages", lambda a: a.transcript_pages or ""), ("Invoice No.", lambda a: a.invoice_no),
    ("File", lambda a: a.file_path),
]
INVOICE_COLUMNS = [h for h, _ in INVOICE_TABLE]
ACTIVITY_COLUMNS = [h for h, _ in ACTIVITY_TABLE]


def invoice_row(i: Invoice) -> list:
    """One invoice as a row of the exported tables (see INVOICE_TABLE)."""
    return [get(i) for _, get in INVOICE_TABLE]


def activity_row(a: Activity) -> list:
    """One activity as a row of the exported tables (see ACTIVITY_TABLE)."""
    return [get(a) for _, get in ACTIVITY_TABLE]


def _from_row(cls, row: sqlite3.Row, **override):
    """A dataclass from a database row: NULL becomes the field's default, unknown columns are dropped (the
    invoices' "idx" column is Invoice.index: INDEX is a word of SQL)."""
    names = {f.name: f for f in fields(cls)}
    raw = dict(row)
    if "idx" in raw:
        raw["index"] = raw.pop("idx")
    d = {k: v for k, v in raw.items() if k in names and v is not None}
    return cls(**{**d, **override})


class Ledger:
    """The records database. mirror_dir: where invoices.csv / activity.csv are kept up to date (None = not)."""

    def __init__(self, path: Path | None = None, mirror_dir: Path | None = None):
        self.path = Path(path) if path else default_db()
        self.mirror_dir = mirror_dir
        self.path.parent.mkdir(parents=True, exist_ok=True)
        with _SCHEMA_LOCK, closing(self._db()) as db:
            db.executescript(_SCHEMA)
            self._migrate(db)
        self.purge()

    @staticmethod
    def _migrate(db: sqlite3.Connection) -> None:
        """Brings a database made by an older version up to date: the columns added since (_ADDED), and every
        invoice number already given entered in used_numbers. Old rows keep blanks in the new columns."""
        with db:
            for table, col, kind in _ADDED:
                have = {r[1] for r in db.execute(f"PRAGMA table_info({table})")}
                if col not in have:
                    try:
                        db.execute(f"ALTER TABLE {table} ADD COLUMN {col} {kind}")
                    except sqlite3.OperationalError as e:  # another program (a second window) added it just now
                        if "duplicate column name" not in str(e):
                            raise
            if db.execute("PRAGMA user_version").fetchone()[0] < SCHEMA_VERSION:
                db.execute("INSERT OR IGNORE INTO used_numbers SELECT invoice_no, year, seq FROM invoices")
                db.execute(f"PRAGMA user_version = {SCHEMA_VERSION}")

    def _db(self) -> sqlite3.Connection:
        db = sqlite3.connect(self.path)
        db.row_factory = sqlite3.Row
        return db

    def _run(self, sql: str, args: tuple = ()) -> None:
        """Makes one change, then brings the CSV copies up to date."""
        db = self._db()
        try:
            with db:
                db.execute(sql, args)
        finally:
            db.close()
        self.mirror()

    def _run_all(self, steps: list[tuple[str, tuple]]) -> None:
        """Makes several changes as one (all or none), then brings the CSV copies up to date."""
        db = self._db()
        try:
            with db:
                for sql, args in steps:
                    db.execute(sql, args)
        finally:
            db.close()
        self.mirror()

    @staticmethod
    def _insert_sql(table: str, values: dict) -> tuple[str, tuple]:
        """The INSERT for one row (Invoice.index is stored in the "idx" column)."""
        values = {("idx" if k == "index" else k): v for k, v in values.items()}
        cols = ", ".join(values)
        return f"INSERT INTO {table} ({cols}) VALUES ({', '.join('?' * len(values))})", tuple(values.values())

    def _insert(self, table: str, values: dict) -> None:
        self._run(*self._insert_sql(table, values))

    def _next_seq(self, year: int) -> int:
        """The next count of the year, after every number given (deleted invoices' too)."""
        rows = self._rows("SELECT MAX(seq) FROM invoices WHERE year = ? UNION ALL "
                          "SELECT MAX(seq) FROM used_numbers WHERE year = ?", (year, year))
        return max([r[0] or 0 for r in rows] + [0]) + 1

    def _set_status(self, invoice_no: str, status: str, speed: str = "", amount: str = "", paid: str = "") -> None:
        self._run("UPDATE invoices SET status=?, paid_speed=?, amount_paid=?, paid_date=? WHERE invoice_no=?",
                  (status, speed, amount, paid, invoice_no))

    def _rows(self, sql: str, args: tuple = ()) -> list[sqlite3.Row]:
        db = self._db()
        try:
            return db.execute(sql, args).fetchall()
        finally:
            db.close()

    # ----------------------------------------------------------- invoice numbers
    def next_invoice_no(self, pattern: str = "{year}-{seq:04}", today: date | None = None) -> tuple[str, int, int]:
        """(number, year, seq): the next number for this year, e.g. '2026-0007'. pattern: the number's format
        ({year}, {yy}, {seq}); a number already taken is skipped. Hold NUMBER_LOCK until add_invoice, or two
        invoices may get the same number."""
        year = (today or date.today()).year
        seq = self._next_seq(year)
        taken = {r[0] for r in self._rows("SELECT invoice_no FROM invoices UNION SELECT invoice_no FROM used_numbers")}

        def number(pattern: str, seq: int) -> str:
            try:
                return pattern.format(year=year, seq=seq, yy=year % 100).strip() or f"{year}-{seq:04}"
            except Exception:  # a bad custom pattern
                return f"{year}-{seq:04}"

        if number(pattern, 1) == number(pattern, 2):  # no {seq}: every invoice would get the same number
            pattern += "-{seq:04}"
        while True:
            no = number(pattern, seq)
            if no not in taken:
                return no, year, seq
            seq += 1

    # ------------------------------------------------------------------ writing
    def log_activity(self, kind: str, **details) -> None:
        """One file made; details are Activity fields (case_name=..., attorney=..., file_path=...)."""
        a = Activity(kind, **details)
        a.ts = a.ts or datetime.now().isoformat(timespec="seconds")
        values = asdict(a)
        values.pop("id")
        values["deleted"] = values["deleted"] or None
        self._insert("activity", values)

    def add_invoice(self, inv: Invoice, year: int = 0, seq: int = 0) -> None:
        """Enters a new invoice. year/seq: from next_invoice_no; without them the invoice counts as the next
        one of its year."""
        year = year or int(inv.created[:4])
        values = asdict(inv)
        values.update(amounts=json.dumps(inv.amounts), year=year, seq=seq or self._next_seq(year))
        values["deleted"] = values["deleted"] or None
        self._run_all([self._insert_sql("invoices", values),
                       ("INSERT OR IGNORE INTO used_numbers VALUES (?, ?, ?)",
                        (inv.invoice_no, year, values["seq"]))])

    def mark_paid(self, invoice_no: str, speed: str, amount, paid_date: str | None = None) -> None:
        """Paid at this speed. amount: '$63.00' or 63; paid_date: ISO ('2026-09-30'), today when not given."""
        self._set_status(invoice_no, "paid", speed, str(money(amount)), paid_date or date.today().isoformat())

    def mark_unpaid(self, invoice_no: str) -> None:
        """Back to open; the payment details are cleared."""
        self._set_status(invoice_no, "open")

    def void(self, invoice_no: str) -> None:
        """Cancelled: it counts 0 from now on (the number stays taken)."""
        self._set_status(invoice_no, "void")

    def set_notes(self, invoice_no: str, notes: str) -> None:
        """Replaces the invoice's notes."""
        self._run("UPDATE invoices SET notes=? WHERE invoice_no=?", (notes, invoice_no))

    def set_amounts(self, invoice_no: str, amounts: dict) -> None:
        """Replaces what each speed of the invoice costs ({"Regular": "$63.00"}), for when an amount was
        corrected in the invoice's PDF: the records would still show the one first made. Each amount is
        kept as "63.00"; ValueError when one is blank or not an amount ("-5", "63 or 70": see rates.parse_amount;
        nothing is changed then). A payment already entered is left as it was."""
        clean = {}
        for speed, value in amounts.items():
            d = parse_amount(value)
            if d is None:
                raise ValueError(f"the amount for {speed or 'a speed'} is not an amount: {value!r}")
            clean[speed] = str(d)
        self._run("UPDATE invoices SET amounts=? WHERE invoice_no=?", (json.dumps(clean), invoice_no))

    # ------------------------------------------------------------------ the trash
    def delete(self, invoices: list[str] = (), activity: list[int] = ()) -> None:
        """Moves records to the trash: invoices by number (with the "invoice" rows of everything made that
        name them) and rows of everything made by id. They can be restored for TRASH_DAYS days."""
        now = datetime.now().isoformat(timespec="seconds")
        steps = []
        for no in invoices:
            steps.append(("UPDATE invoices SET deleted=? WHERE invoice_no=? AND deleted IS NULL", (now, no)))
            steps.append(("UPDATE activity SET deleted=? WHERE invoice_no=? AND kind='invoice' AND deleted IS NULL",
                          (now, no)))
        steps += [("UPDATE activity SET deleted=? WHERE id=? AND deleted IS NULL", (now, i)) for i in activity]
        if steps:
            self._run_all(steps)

    def restore(self, invoices: list[str] = (), activity: list[int] = ()) -> None:
        """Takes records out of the trash (an invoice with the rows of everything made that name it)."""
        steps = []
        for no in invoices:
            steps.append(("UPDATE invoices SET deleted=NULL WHERE invoice_no=?", (no,)))
            steps.append(("UPDATE activity SET deleted=NULL WHERE invoice_no=? AND kind='invoice'", (no,)))
        steps += [("UPDATE activity SET deleted=NULL WHERE id=?", (i,)) for i in activity]
        if steps:
            self._run_all(steps)

    def delete_forever(self, invoices: list[str] = (), activity: list[int] = ()) -> None:
        """Deletes records in the trash for good (records not in the trash are left alone). An invoice's
        number stays taken."""
        steps = []
        for no in invoices:
            steps.append(("DELETE FROM invoices WHERE invoice_no=? AND deleted IS NOT NULL", (no,)))
            steps.append(("DELETE FROM activity WHERE invoice_no=? AND kind='invoice' AND deleted IS NOT NULL", (no,)))
        steps += [("DELETE FROM activity WHERE id=? AND deleted IS NOT NULL", (i,)) for i in activity]
        if steps:
            self._run_all(steps)

    def purge(self, now: datetime | None = None) -> int:
        """Deletes for good what has been in the trash longer than TRASH_DAYS days; returns how many records.
        Done each time the records are opened."""
        cutoff = ((now or datetime.now()) - timedelta(days=TRASH_DAYS)).isoformat(timespec="seconds")
        db = self._db()
        try:
            with db:
                n = db.execute("DELETE FROM invoices WHERE deleted IS NOT NULL AND deleted < ?", (cutoff,)).rowcount
                n += db.execute("DELETE FROM activity WHERE deleted IS NOT NULL AND deleted < ?", (cutoff,)).rowcount
        finally:
            db.close()
        if n:
            self.mirror()
        return n

    # ------------------------------------------------------------------ reading
    def invoices(self, year: int | None = None, month: int | None = None, client: str = "",
                 status: str = "", text: str = "", trash: bool = False, how: str = "words") -> list[Invoice]:
        """Newest first, filtered by year, month, status ("open", "paid", "void"), client (the exact
        firm/attorney, see Invoice.client) and text: found anywhere in the invoice (number, case, index number,
        firm, attorney, e-mail, court, part, judge, dates, excerpt, reporters, notes, the speed paid, file),
        as `how` says
        (see matcher: "words", "regex" or "fuzzy"; ValueError for a pattern that isn't one). trash: the
        invoices in the trash instead of the others."""
        found = matcher(text, how)
        out = []
        where = "deleted IS NOT NULL" if trash else "deleted IS NULL"
        for r in self._rows(f"SELECT * FROM invoices WHERE {where} "
                            "ORDER BY created DESC, year DESC, seq DESC, invoice_no DESC"):
            try:
                amounts = json.loads(r["amounts"] or "{}")
            except ValueError:
                amounts = {}
            inv = _from_row(Invoice, r, amounts=amounts)
            if year and inv.created[:4] != str(year):
                continue
            if month and inv.created[5:7] != f"{month:02}":
                continue
            if client and inv.client != client:
                continue
            if status and inv.status != status:
                continue
            if not found([inv.invoice_no, inv.case_name, inv.index_no, inv.bill_to, inv.firm, inv.email,
                          inv.court, inv.part, inv.judge, inv.dates, inv.excerpt, inv.reporters, inv.notes,
                          inv.paid_speed, Path(inv.file_path).name if inv.file_path else ""]):
                continue
            out.append(inv)
        return out

    def invoice(self, invoice_no: str) -> Invoice | None:
        """One invoice by its number (in the trash or not), or None."""
        return next((i for i in self.invoices() + self.invoices(trash=True) if i.invoice_no == invoice_no), None)

    def activity(self, kind: str = "", since: str = "", until: str = "", text: str = "",
                 trash: bool = False, how: str = "words") -> list[Activity]:
        """Newest first. kind: a key of KINDS; since/until: ISO dates (inclusive); text: found anywhere in the
        row (case, index number, attorney, firm, judge, part, dates, invoice number, what was made as the
        Activity tab names it ("Minute agreement"), file), as `how` says (see matcher); trash: the rows in the
        trash instead of the others."""
        found = matcher(text, how)
        out = []
        where = "deleted IS NOT NULL" if trash else "deleted IS NULL"
        for r in self._rows(f"SELECT * FROM activity WHERE {where} ORDER BY ts DESC, id DESC"):
            a = _from_row(Activity, r)
            if kind and a.kind != kind:
                continue
            if since and a.ts[:10] < since:
                continue
            if until and a.ts[:10] > until:
                continue
            if not found([a.case_name, a.index_no, a.attorney, a.firm, a.judge, a.part, a.dates, a.invoice_no,
                          KINDS.get(a.kind, a.kind), Path(a.file_path).name if a.file_path else ""]):
                continue
            out.append(a)
        return out

    def years(self) -> list[int]:
        """The years with any invoice or activity, newest first."""
        rows = self._rows("SELECT DISTINCT substr(created,1,4) FROM invoices WHERE deleted IS NULL UNION "
                          "SELECT DISTINCT substr(ts,1,4) FROM activity WHERE deleted IS NULL")
        return sorted({int(r[0]) for r in rows if r[0] and r[0].isdigit()}, reverse=True)

    def clients(self) -> list[str]:
        """Every firm or attorney billed (see Invoice.client), A to Z."""
        return sorted({i.client for i in self.invoices()}, key=str.lower)

    # ------------------------------------------------------------------ exports
    def mirror(self) -> None:
        """Keeps the CSV copies in the records folder current; a file open in Excel is skipped (and logged)."""
        if self.mirror_dir is None:
            return
        try:
            with _MIRROR_LOCK:
                self.export_csv(self.mirror_dir)
        except Exception as e:  # the change itself is saved; the copies catch up with the next one
            log_error("could not update the CSV copies of the records", e)

    def export_csv(self, folder: Path) -> list[Path]:
        """Writes invoices.csv and activity.csv (every row not in the trash) into folder; returns their paths."""
        folder = Path(folder)
        folder.mkdir(parents=True, exist_ok=True)
        out = []
        for name, cols, rows in (("invoices.csv", INVOICE_COLUMNS, [invoice_row(i) for i in self.invoices()]),
                                 ("activity.csv", ACTIVITY_COLUMNS, [activity_row(a) for a in self.activity()])):
            p = folder / name
            with open(p, "w", newline="", encoding="utf-8-sig") as f:  # the BOM makes Excel read it as UTF-8
                w = csv.writer(f)
                w.writerow(cols)
                w.writerows(rows)
            out.append(p)
        return out

    def export_xlsx(self, path: Path, invoices: list[Invoice] | None = None) -> Path:
        """An Excel workbook: Invoices, Activity, By firm and By month sheets. invoices: the ones to list
        (default: all); the Activity sheet always has every row."""
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter

        invoices = self.invoices() if invoices is None else invoices
        total, by_client, by_month = summarize(invoices)
        wb = Workbook()
        head = Font(bold=True, color="FFFFFF")
        fill = PatternFill("solid", fgColor="4B3F8F")

        def sheet(ws, cols, rows, money_cols=()):
            ws.append(cols)
            for c in ws[1]:
                c.font, c.fill = head, fill
            for r in rows:
                ws.append(r)
            for i, col in enumerate(cols, 1):
                letter = get_column_letter(i)
                width = max([len(str(col))] + [len(str(r[i - 1])) for r in rows[:500]])
                ws.column_dimensions[letter].width = min(60, width + 2)
                if col in money_cols:
                    for cell in ws[letter][1:]:
                        cell.number_format = '"$"#,##0.00'
            ws.freeze_panes = "A2"
            if rows:
                ws.auto_filter.ref = ws.dimensions

        ws = wb.active
        ws.title = "Invoices"
        sheet(ws, INVOICE_COLUMNS, [invoice_row(i) for i in invoices], ("Billed", "Amount paid"))
        sheet(wb.create_sheet("Activity"), ACTIVITY_COLUMNS, [activity_row(a) for a in self.activity()])
        sums = ["Invoices", "Billed", "Paid", "Outstanding"]
        sheet(wb.create_sheet("By firm"), ["Firm / attorney"] + sums,
              [[k, v.count, float(v.billed), float(v.paid), float(v.outstanding)] for k, v in by_client.items()]
              + [["Total", total.count, float(total.billed), float(total.paid), float(total.outstanding)]],
              ("Billed", "Paid", "Outstanding"))
        sheet(wb.create_sheet("By month"), ["Month"] + sums,
              [[_month_name(k), v.count, float(v.billed), float(v.paid), float(v.outstanding)]
               for k, v in by_month.items()], ("Billed", "Paid", "Outstanding"))
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)
        return path

    def export_html(self, path: Path, invoices: list[Invoice] | None = None, title: str = "Invoice report") -> Path:
        """A self-contained report page (totals, a monthly chart, per-firm and per-invoice tables)."""
        invoices = self.invoices() if invoices is None else invoices
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(report_html(invoices, title), encoding="utf-8")
        return path


SEARCH_MODES = ("words", "regex", "fuzzy")
FUZZY_SCORE = 80  # how alike a word must be (0-100, rapidfuzz's partial_ratio) to count as found


def matcher(text: str, how: str = "words"):
    """A test for one record's values (a list of strings) against the search text:
      "words"  every word of text is somewhere in them, any case ("counsel 2026"); a dash and a slash find
               each other, both ways ("712345-2021" finds "712345/2021", "6/3/2026" finds "6-3-2026")
      "regex"  text is a regular expression found somewhere in them, any case ("^Invoice 2026-00(1|2)");
               ValueError when it isn't a valid one
      "fuzzy"  every word is close to a part of them (rapidfuzz's partial_ratio at least FUZZY_SCORE), so a
               misspelling still finds it ("Counsle" finds "Counsel & Counsel"); words of one or two letters
               must be found as they are, and in practice so must words of three or four (one wrong letter
               already puts them below the score); ValueError when rapidfuzz is missing
    Blank text finds everything. ValueError for a `how` not in SEARCH_MODES."""
    if how not in SEARCH_MODES:
        raise ValueError(f"unknown search mode {how!r} (one of {', '.join(SEARCH_MODES)})")
    text = text.strip()
    if not text:
        return lambda values: True
    if how == "regex":
        import re
        try:
            pattern = re.compile(text, re.IGNORECASE)
        except re.error as e:
            raise ValueError(f"not a valid regular expression: {e}") from None
        return lambda values: any(pattern.search(str(v)) for v in values if v)
    words = text.lower().split()
    if how == "fuzzy":
        try:
            from rapidfuzz import fuzz
        except ImportError:  # (a build without it): the box says so, as for a pattern that can't be read
            raise ValueError("Fuzzy search isn't available: the rapidfuzz library is missing") from None

        def close(values) -> bool:
            hay = " ".join(str(v) for v in values if v).lower()
            return all(w in hay or (len(w) > 2 and fuzz.partial_ratio(w, hay) >= FUZZY_SCORE) for w in words)
        return close

    def has(values) -> bool:
        hay = " ".join(str(v) for v in values if v).lower()
        return all(w in hay or w.replace("-", "/") in hay or w.replace("/", "-") in hay for w in words)
    return has


def _month_name(ym: str) -> str:
    """'2026-09' -> 'Sep 2026'."""
    try:
        return date(int(ym[:4]), int(ym[5:7]), 1).strftime("%b %Y")
    except ValueError:
        return ym


def report_html(invoices: list[Invoice], title: str = "Invoice report") -> str:
    """The HTML report as one string: no outside files, works offline, light and dark."""
    e = html.escape
    total, by_client, by_month = summarize(invoices)
    tiles = "".join(f'<div class="tile"><div class="k">{k}</div><div class="v">{v}</div></div>' for k, v in (
        ("Billed", fmt(total.billed)), ("Paid", fmt(total.paid)), ("Outstanding", fmt(total.outstanding)),
        ("Invoices", str(total.count))))

    # monthly bars: paid (solid) stacked under outstanding (lighter)
    months = list(by_month.items())[-24:]
    top = max([float(v.billed) for _, v in months] + [1.0])
    bw, gap, h = 28, 10, 160
    bars = []
    for i, (m, v) in enumerate(months):
        x = 40 + i * (bw + gap)
        ph = float(v.paid) / top * h
        oh = float(v.outstanding) / top * h
        bars.append(f'<g><title>{e(_month_name(m))}: billed {fmt(v.billed)}, paid {fmt(v.paid)}</title>'
                    f'<rect x="{x}" y="{20 + h - ph:.1f}" width="{bw}" height="{ph:.1f}" class="paid"/>'
                    f'<rect x="{x}" y="{20 + h - ph - oh:.1f}" width="{bw}" height="{oh:.1f}" class="due"/>'
                    f'<text x="{x + bw / 2}" y="{h + 36}" class="lbl">{e(_month_name(m)[:3])}</text>'
                    f'<text x="{x + bw / 2}" y="{h + 50}" class="lbl">{m[2:4]}</text></g>')
    width = 40 + max(1, len(months)) * (bw + gap)
    chart = (f'<svg viewBox="0 0 {width} {h + 60}" role="img" aria-label="Billed per month">'
             f'<line x1="36" y1="{20 + h}" x2="{width}" y2="{20 + h}" class="axis"/>'
             f'<text x="34" y="26" class="lbl end">{fmt(Decimal(top))}</text>{"".join(bars)}</svg>'
             if months else "<p class='muted'>No invoices yet.</p>")

    def rows(cells: list[list[str]]) -> str:
        return "".join("<tr>" + "".join(f"<td>{c}</td>" for c in r) + "</tr>" for r in cells)

    firm_rows = rows([[e(k), str(v.count), fmt(v.billed), fmt(v.paid), fmt(v.outstanding)]
                      for k, v in by_client.items()])
    month_rows = rows([[e(_month_name(k)), str(v.count), fmt(v.billed), fmt(v.paid), fmt(v.outstanding)]
                       for k, v in reversed(list(by_month.items()))])
    inv_rows = rows([[e(i.invoice_no), us_date(i.created), e(i.case_name), e(i.index_no), e(i.client),
                      str(i.pages), fmt(i.billed),
                      f'<span class="st {i.status}">{e(i.status.title())}</span>'
                      + (f' <span class="muted">{e(i.paid_speed)} {us_date(i.paid_date)}</span>' if i.status == "paid"
                         else "")]
                     for i in invoices])
    stamp = datetime.now().strftime("%B %d, %Y %I:%M %p").replace(" 0", " ")
    return f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>{e(title)}</title>
<style>
:root {{ --bg:#f7f6fb; --card:#fff; --ink:#1d1b26; --muted:#6b6880; --line:#e4e1ee; --accent:#5b4bc4;
        --accent2:#c9c2f0; --ok:#1f8a4c; --due:#b5651d; --void:#8a8797; }}
@media (prefers-color-scheme: dark) {{ :root {{ --bg:#15141b; --card:#1f1d28; --ink:#ecebf3; --muted:#a19eb3;
        --line:#2f2c3b; --accent:#9d8ff0; --accent2:#4a4278; --ok:#5fcf8f; --due:#e3a35c; --void:#8a8797; }} }}
* {{ box-sizing:border-box; }}
body {{ margin:0; background:var(--bg); color:var(--ink); font:15px/1.45 "Segoe UI", system-ui, sans-serif; }}
main {{ max-width:1100px; margin:0 auto; padding:24px 16px 48px; }}
h1 {{ font-size:26px; margin:0 0 4px; }} h2 {{ font-size:17px; margin:28px 0 10px; }}
.muted {{ color:var(--muted); }}
.tiles {{ display:grid; grid-template-columns:repeat(auto-fit, minmax(170px, 1fr)); gap:12px; margin-top:18px; }}
.tile {{ background:var(--card); border:1px solid var(--line); border-radius:12px; padding:14px 16px; }}
.tile .k {{ color:var(--muted); font-size:13px; }} .tile .v {{ font-size:24px; font-weight:600; }}
.card {{ background:var(--card); border:1px solid var(--line); border-radius:12px; padding:14px; overflow-x:auto; }}
svg {{ width:100%; max-height:260px; }} .paid {{ fill:var(--accent); }} .due {{ fill:var(--accent2); }}
.axis {{ stroke:var(--line); }} .lbl {{ fill:var(--muted); font-size:10px; text-anchor:middle; }} .end {{ text-anchor:end; }}
table {{ width:100%; border-collapse:collapse; font-size:14px; }}
th, td {{ text-align:left; padding:7px 8px; border-bottom:1px solid var(--line); vertical-align:top; }}
th {{ color:var(--muted); font-weight:600; font-size:12px; text-transform:uppercase; letter-spacing:.03em; }}
.st {{ font-weight:600; }} .st.paid {{ color:var(--ok); }} .st.open {{ color:var(--due); }} .st.void {{ color:var(--void); }}
.legend span {{ display:inline-block; width:10px; height:10px; border-radius:2px; margin:0 4px 0 12px; }}
</style></head><body><main>
<h1>{e(title)}</h1><div class="muted">Made by DjinnIt on {e(stamp)}</div>
<div class="tiles">{tiles}</div>
<h2>Billed per month</h2>
<div class="card">{chart}<div class="muted legend"><span style="background:var(--accent)"></span>Paid
<span style="background:var(--accent2)"></span>Outstanding</div></div>
<h2>By firm</h2><div class="card"><table><tr><th>Firm / attorney</th><th>Invoices</th><th>Billed</th><th>Paid</th>
<th>Outstanding</th></tr>{firm_rows}</table></div>
<h2>By month</h2><div class="card"><table><tr><th>Month</th><th>Invoices</th><th>Billed</th><th>Paid</th>
<th>Outstanding</th></tr>{month_rows}</table></div>
<h2>Invoices</h2><div class="card"><table><tr><th>No.</th><th>Date</th><th>Case</th><th>Index No.</th>
<th>Bill to</th><th>Pages</th><th>Billed</th><th>Status</th></tr>{inv_rows}</table></div>
</main></body></html>
"""
