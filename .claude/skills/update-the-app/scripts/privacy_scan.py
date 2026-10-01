"""Checks what is about to be committed for real names, numbers and contact details.

The words to look for come from the git-ignored samples_internal/ folder, so no real name ever has to be
written in the repository:
  - expected.json: "_profile" (the user's name, e-mail, phone), the names and numbers in the file names
    of the samples, and every value expected from them (judges, attorneys, firms, index numbers)
  - private_words.txt: anything else, one per line (colleagues, witnesses, parties); # starts a comment

Run from anywhere in the repository after `git add`:
    .venv/Scripts/python.exe .claude/skills/update-the-app/scripts/privacy_scan.py [--worktree]
--worktree scans every change not yet committed (staged or not) and new files instead of the staged ones.
The names of the files are checked too, and the text of Excel workbooks and PDFs being added.
Exit code 1 when something is found (or the check can't be done). Pictures can't be read: they are listed,
look at them yourself.
"""
from __future__ import annotations

import io
import json
import re
import subprocess
import sys
from pathlib import Path


def _git(*args: str, cwd: Path | None = None, binary: bool = False):
    """A git command's output; exits (code 1) when git fails."""
    r = subprocess.run(["git", *args], cwd=cwd, capture_output=True)
    if r.returncode != 0:
        sys.exit(f"git {' '.join(args)} failed: {r.stderr.decode('utf-8', 'replace').strip()}")
    return r.stdout if binary else r.stdout.decode("utf-8", "replace")


try:
    ROOT = Path(_git("rev-parse", "--show-toplevel", cwd=Path(__file__).resolve().parent).strip())
except FileNotFoundError:
    sys.exit("git was not found")
PRIVATE = ROOT / "samples_internal"
# Words of real file names and values that say nothing about anyone
COMMON = {"transcript", "copy", "invoice", "minutes", "page", "sheet", "supreme", "queens", "court", "county", "york",
          "civil", "trial", "plaintiff", "defendant", "group", "office", "offices", "firm", "law", "esq", "llp", "pllc",
          "the", "and", "for", "state", "city", "health", "hospital", "hospitals", "care", "center", "medical",
          "associates", "attorneys", "private", "rates", "choice", "form", "agreement", "fillable", "minute", "auto",
          "billing", "corporation", "company"}


def terms() -> tuple[set[str], set[str]]:
    """(words and phrases, phone numbers as digits) to look for."""
    words: set[str] = set()
    phones: set[str] = set()

    def person(name: str, every_word: bool = False) -> None:
        """A name, and its last word (a surname: first names like "John" would match made-up names too).
        every_word: also each word (the user's own name)."""
        name = " ".join(str(name).split())
        if len(name) >= 4 and name.lower() not in COMMON:
            words.add(name)
        parts = [w for w in re.findall(r"[A-Za-z][A-Za-z'\-]{3,}", name) if w.lower() not in COMMON]
        words.update(parts if every_word else parts[-1:])

    exp = PRIVATE / "expected.json"
    if exp.exists():
        data = json.loads(exp.read_text(encoding="utf-8"))
        for key, value in (data.get("_profile") or {}).items():
            if key == "phone":
                phones.add(re.sub(r"\D", "", value)[-10:])
            elif key == "email":
                words.add(value)
                words.add(value.split("@")[0])
            else:
                person(value, every_word=True)
        for name, entry in data.items():
            if name.startswith("_"):
                continue
            stem = Path(name).stem
            words.update(n for n in re.findall(r"\d{5,7}", stem))
            for w in re.findall(r"[A-Za-z][A-Za-z'\-]{3,}", stem):
                if w.lower() not in COMMON and not re.fullmatch(r"PXL|IMG", w):
                    words.add(w)
            if not isinstance(entry, dict):
                continue
            for group in ("fields", "fields_startswith"):
                for k, v in (entry.get(group) or {}).items():
                    if k in ("judge", "case_name"):
                        person(v)
                    elif k == "index_no":
                        words.add(re.split(r"\D", v)[0])
            for a in entry.get("attorneys") or []:
                person(a)
            for a, extra in (entry.get("attorney_fields") or {}).items():
                person(a)
                for v in extra.values():
                    person(v)
    extra = PRIVATE / "private_words.txt"
    if extra.exists():
        for line in extra.read_text(encoding="utf-8").splitlines():
            line = line.split("#")[0].strip()
            digits = re.sub(r"\D", "", line)
            if line and not re.search(r"[A-Za-z]", line) and len(digits) >= 10:  # a phone number
                phones.add(digits[-10:])
            elif line:
                words.add(line)
    return words, phones


def document_text(name: str, data: bytes) -> str | None:
    """The text of an Excel workbook or a PDF (every cell, every page); None for other files."""
    suffix = Path(name).suffix.lower()
    try:
        if suffix in (".xlsx", ".xlsm"):
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(data), read_only=True)
            try:
                return "\n".join(" | ".join(str(c) for c in row if c is not None)
                                 for ws in wb.worksheets for row in ws.iter_rows(values_only=True)) + \
                    "\n" + "\n".join(wb.sheetnames)
            finally:
                wb.close()
        if suffix == ".pdf":
            import pymupdf
            with pymupdf.open(stream=data, filetype="pdf") as doc:
                return "\n".join(page.get_text() for page in doc) + "\n" + \
                    "\n".join(str(v) for v in (doc.metadata or {}).values() if v)
    except Exception as e:
        return f"(could not be read: {type(e).__name__})"
    return None


def changes(worktree: bool) -> tuple[list[tuple[str, str]], list[str], list[str]]:
    """(file, line) for every line added by the staged changes (or by every uncommitted change and the new
    files), the paths of the files changed or added, and the pictures among them."""
    diff = ["diff", "HEAD"] if worktree else ["diff", "--cached"]
    out = _git(*diff, "-U0", "--no-color", "--no-ext-diff", cwd=ROOT)
    lines, current = [], ""
    for line in out.splitlines():
        if line.startswith("+++ "):
            current = line[6:] if line.startswith("+++ b/") else line[4:]
        elif line.startswith("+") and not line.startswith("+++"):
            lines.append((current, line[1:]))
    paths = [p for p in _git(*diff, "--name-only", "--diff-filter=ACMR", cwd=ROOT).splitlines() if p]
    binary = [l.split("\t")[-1] for l in _git(*diff, "--numstat", cwd=ROOT).splitlines() if l.startswith("-\t-\t")]
    new = [p for p in _git("ls-files", "--others", "--exclude-standard", cwd=ROOT).splitlines() if p] \
        if worktree else []
    for f in new:  # new files: every line
        try:
            data = (ROOT / f).read_bytes()
        except OSError:
            continue
        if b"\0" in data:
            binary.append(f)
            continue
        lines += [(f, l) for l in data.decode("utf-8", "replace").splitlines()]
    for f in dict.fromkeys(binary):  # workbooks and PDFs: their text
        try:
            data = (ROOT / f).read_bytes() if worktree else _git("show", f":{f}", cwd=ROOT, binary=True)
        except (OSError, SystemExit):
            continue
        text = document_text(f, data)
        if text is not None:
            lines += [(f, l) for l in text.splitlines()]
    # pictures and other files that can't be read here
    unread = [f for f in dict.fromkeys(binary) if Path(f).suffix.lower() not in (".xlsx", ".xlsm", ".pdf")]
    return lines, list(dict.fromkeys(paths + new)), unread


def main() -> int:
    if not PRIVATE.is_dir():
        print("samples_internal/ is missing: there is nothing to compare against - the check can't be done.")
        return 1
    ignored = subprocess.run(["git", "check-ignore", "-q", "samples_internal/x"], cwd=ROOT).returncode == 0
    if not ignored:
        print("DANGER: samples_internal/ is not git-ignored!")
        return 1
    words, phones = terms()
    patterns = [(w, re.compile(r"(?<![A-Za-z0-9])" + re.escape(w) + r"(?![A-Za-z0-9])", re.I))
                for w in sorted(words, key=len, reverse=True)]

    def found_in(text: str) -> list[str]:
        digits = re.sub(r"\D", "", text)
        return [w for w, p in patterns if p.search(text)] + [p for p in phones if p and p in digits]

    lines, paths, pictures = changes("--worktree" in sys.argv)
    hits = []
    for f in paths:
        if f.startswith("samples_internal/"):
            hits.append((f, "(a private file is to be committed!)", ""))
        elif found_in(f):
            hits.append((f, ", ".join(found_in(f)[:4]), "(in the file's name)"))
    for f, line in lines:
        if f.startswith("samples_internal/"):
            continue
        found = found_in(line)
        if found:
            hits.append((f, ", ".join(found[:4]), line.strip()[:140]))
    print(f"checked {len(paths)} file(s) against {len(words)} private words and {len(phones)} phone number(s)")
    for f, what, line in hits:
        print(f"  {f}: [{what}]  {line}")
    if pictures:
        print("look at these yourself (made-up data only):", ", ".join(pictures))
    print("FOUND private data - replace it with made-up data" if hits else "clean")
    return 1 if hits else 0


if __name__ == "__main__":
    sys.exit(main())
