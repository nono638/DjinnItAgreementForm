"""Runs the built program headless on made-up inputs and checks what it made.

After PyInstaller (from anywhere; PYTHONPATH is not needed):
    .venv/Scripts/python.exe .claude/skills/update-the-app/scripts/exe_check.py

It uses a temporary APPDATA and home folder, so the user's own settings, records and run sheets are not
touched (the folder is removed again when all is good); its settings say the user is Pat Reporter (initials
P.R.). The inputs are two days of one trial (a transcript each, June 3 and June 4, written by Pat and Dana
Smith: PR and DS at the foot of the pages) and an e-mail about another case. Expected: the built version, no
errors in the log; for each day a minute agreement, a MOFR and the run sheet, and one joint invoice for both
days (listed under each day) that bills Pat's own 10 of the 20 transcript pages, for one party, and has its
amounts as fields (and no "Lock fields" button, removed in 1.3.1); one run sheet with the 4 takes of each day;
and --selftest says the libraries only the window uses are in the build (the Records' Fuzzy and Regex
searches, HEIC photos, the look for a newer version, printing, The math as a PDF). Exit code 1 if something is
off.
"""
import csv
import json
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[4]  # .claude/skills/update-the-app/scripts/ -> the repository
sys.path[:0] = [str(ROOT), str(ROOT / "tests")]
import pymupdf  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from minute_filler import __version__  # noqa: E402
from test_runsheet import transcript  # noqa: E402

sys.stdout.reconfigure(encoding="utf-8", errors="replace")  # (the report may hold arrows and ellipses)


def kind(name: str) -> str:
    """'Invoice 2026-0001 - Jane Roe v. Sam Poe.pdf' -> 'Invoice'; the Excel file is the run sheet."""
    return "Run sheet" if name.endswith(".xlsx") else name.split(" - ")[0].split(" 20")[0]


exe = ROOT / "dist" / "DjinnItAgreementForm" / "DjinnItAgreementForm.exe"
if not exe.exists():
    sys.exit(f"PROBLEM: {exe} is not there - build it first (PyInstaller)")
run = Path(tempfile.mkdtemp(prefix="djinnit-exe-check-"))
src = run / "in"
src.mkdir()
transcript(src / "Transcript 6-3-2026 Roe v Poe.pdf")  # two reporters, a witness, a word index at the end
transcript(src / "Transcript 6-4-2026 Roe v Poe.pdf", day="June 4, 2026")  # the next day: one invoice for both
(src / "email.txt").write_text("From: someone@examplefirm.com\nSubject: minutes\n\nPlease send the minutes in "
                               "Smith v Jones, Index No. 712222/2024, 5/22/2026, Judge Lopez. About 40 pages.")
# who the user is: the transcripts were written by two reporters, and only the user's pages are billed (without
# the initials, the invoice would be held until Whose pages... says)
settings = run / "appdata" / "DjinnItAgreementForm" / "settings.json"
settings.parent.mkdir(parents=True)
settings.write_text(json.dumps({"profile": {"name": "Pat Reporter", "initials": "P.R."}}), encoding="utf-8")
env = dict(os.environ, APPDATA=str(run / "appdata"), USERPROFILE=str(run / "home"), HOME=str(run / "home"))
problems = []
try:
    subprocess.run([str(exe), "--batch", str(run / "out"), "--outputs", "agreement,mofr,invoice,runsheet", str(src)],
                   env=env, timeout=300)  # exit code 1 is expected: the e-mail job has no transcript
except subprocess.TimeoutExpired:
    sys.exit(f"PROBLEM: the program did not finish within 5 minutes (files in {run})")

report_file = run / "out" / "batch.json"
if not report_file.exists():
    sys.exit(f"PROBLEM: the program wrote no batch.json (files in {run})")
report = json.loads(report_file.read_text(encoding="utf-8"))
for j in report["jobs"]:
    print(" ", j["case"], "|", j["forms"], "|", j["error"] or "ok")
days = sorted((j for j in report["jobs"] if "Roe" in (j["case"] or "")), key=lambda j: j["dates"] or "")
kinds = [sorted(kind(f) for f in j["forms"]) for j in days]
if len(days) != 2 or any(j["error"] for j in days) or kinds != [["Invoice", "MOFR", "Minute Agreement",
                                                                  "Run sheet"]] * 2:
    problems.append("each day of the trial should have an agreement, a MOFR, the run sheet and the joint invoice")
invoices = sorted({f for j in days for f in j["forms"] if f.startswith("Invoice")})
if len(invoices) != 1:
    problems.append(f"the two days should share one invoice, not {len(invoices)}")
else:
    with pymupdf.open(run / "out" / invoices[0]) as doc:
        widgets = {w.field_name for w in doc[0].widgets()}
    if not {"amount Regular", "amount Expedite"} <= widgets or "DjinnIt lock" in widgets:
        problems.append("the invoice should have its amounts as fields, and no Lock fields button")
invoices_csv = run / "home" / "Documents" / "DjinnIt Records" / "invoices.csv"
rows = []
if invoices_csv.exists():
    with open(invoices_csv, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
if len(rows) != 1:
    problems.append(f"one invoice should be recorded, not {len(rows)}")
else:
    row = rows[0]
    print("  invoice:", {k: row.get(k) for k in ("Pages", "Transcript pages", "Reporters", "Parties")})
    if (row.get("Pages"), row.get("Transcript pages"), row.get("Parties")) != ("10", "20", "1") \
            or row.get("Date(s) of proceeding") != "6/3/2026, 6/4/2026":
        problems.append("the invoice should bill Pat's 10 of both days' 20 transcript pages (not the word index), "
                        "for one party")
sheets = list((run / "home" / "Documents" / "DjinnIt Run Sheets").glob("*.xlsx"))
if sheets:
    ws = load_workbook(sheets[0])["Run Sheet"]
    takes = [(ws.cell(r, 3).value, ws.cell(r, 6).value) for r in range(5, ws.max_row + 1)]
    print("  run sheet:", takes)
    if len(sheets) != 1 or takes != [("Pat", 2), ("Dana", 3), ("Pat", 3), ("Dana", 2)] * 2:
        problems.append("one run sheet should have the takes of both days")
else:
    problems.append("no run sheet was made")
# --selftest: the libraries the window needs that the batch doesn't use (rapidfuzz, regex, pillow-heif, ssl, Qt print)
try:
    subprocess.run([str(exe), "--selftest", str(run / "st")], env=env, timeout=120)
except subprocess.TimeoutExpired:
    problems.append("--selftest did not finish within 2 minutes")
st_file = run / "st" / "selftest.json"
st = json.loads(st_file.read_text(encoding="utf-8")) if st_file.exists() else {}
for key, what, lib in (("fuzzy_search", "the Records' Fuzzy search", "rapidfuzz"),
                       ("regex_search", "the Records' Regex search", "regex"),
                       ("heic_photos", "reading iPhone HEIC photos", "pillow-heif"),
                       ("update_check", "the look for a newer version", "ssl"),
                       ("printing", "Print…", "PySide6.QtPrintSupport"),
                       ("math_pdf", "Save as PDF in The math", "pymupdf (Story)")):
    got = st.get(key, "no selftest.json" if not st else "missing from selftest.json")
    print(f"  {key.replace('_', ' ')}:", got)
    if got is not True:
        problems.append(f"{what} does not work in the build ({got}): is {lib} bundled?")
log_file = run / "appdata" / "DjinnItAgreementForm" / "logs" / "app.log"
log = log_file.read_text(encoding="utf-8") if log_file.exists() else ""
started = [l for l in log.splitlines() if "started: version" in l]
print("  log:", started[-1][20:80] if started else "(no start line)", "| errors:", log.count("ERROR"))
if not started or f"version {__version__} " not in started[-1]:
    problems.append(f"the built program is not version {__version__} (rebuild?)")
if log.count("ERROR"):
    problems.append("the log has errors")
if problems:
    print("\n".join("PROBLEM: " + p for p in problems) + f"\n(files in {run})")
else:
    print("all good")
    shutil.rmtree(run, ignore_errors=True)
sys.exit(1 if problems else 0)
