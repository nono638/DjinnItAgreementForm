"""The records database: activity log, invoice ledger, totals, paid status and exports."""
import csv
from decimal import Decimal

import pytest

from minute_filler.records import Invoice, Ledger, summarize


def inv(no, created, firm, amounts, **kw):
    return Invoice(invoice_no=no, created=created, case_name=f"Case {no}", index_no="700001/2025", bill_to="Alex",
                   firm=firm, pages=10, parties=1, amounts=amounts, billed_speed=next(iter(amounts)), **kw)


@pytest.fixture
def ledger(tmp_path):
    lg = Ledger(tmp_path / "records.db", mirror_dir=tmp_path / "mirror")
    lg.add_invoice(inv("2026-0001", "2026-01-15", "Example Firm LLP", {"Regular": "63.00", "Daily": "91.00"}))
    lg.add_invoice(inv("2026-0002", "2026-01-20", "Advocate & Partners", {"Regular": "100.00"}))
    lg.add_invoice(inv("2026-0003", "2026-02-02", "Example Firm LLP", {"Regular": "40.00", "Daily": "55.00"}))
    lg.add_invoice(inv("2025-0009", "2025-12-30", "", {"Regular": "10.00"}))
    return lg


def test_numbers_count_up_per_year(ledger):
    from datetime import date
    assert ledger.next_invoice_no(today=date(2026, 5, 1))[0] == "2026-0004"
    assert ledger.next_invoice_no(today=date(2027, 1, 2))[0] == "2027-0001"
    assert ledger.next_invoice_no("INV-{yy}{seq:03}", today=date(2026, 5, 1))[0] == "INV-26004"
    assert ledger.next_invoice_no("{nonsense}", today=date(2026, 5, 1))[0] == "2026-0004"


def test_filters(ledger):
    assert [i.invoice_no for i in ledger.invoices()] == ["2026-0003", "2026-0002", "2026-0001", "2025-0009"]
    assert len(ledger.invoices(year=2026)) == 3 and len(ledger.invoices(year=2026, month=1)) == 2
    assert len(ledger.invoices(client="Example Firm LLP")) == 2
    assert [i.invoice_no for i in ledger.invoices(text="case 2026-0002")] == ["2026-0002"]
    assert ledger.invoices(client="Alex")[0].invoice_no == "2025-0009"  # no firm: grouped by attorney
    assert ledger.years() == [2026, 2025]
    assert ledger.clients() == ["Advocate & Partners", "Alex", "Example Firm LLP"]


def test_paid_at_another_speed_and_void(ledger):
    total, by_client, by_month = summarize(ledger.invoices(year=2026))
    assert (total.count, total.billed, total.paid, total.outstanding) == (3, Decimal("203.00"), 0, Decimal("203.00"))
    ledger.mark_paid("2026-0001", "Daily", "91", "2026-02-01")
    ledger.void("2026-0002")
    i = ledger.invoice("2026-0001")
    assert (i.status, i.paid_speed, i.amount_paid, i.paid_date) == ("paid", "Daily", "91.00", "2026-02-01")
    total, by_client, by_month = summarize(ledger.invoices(year=2026))
    assert (total.count, total.billed, total.paid, total.outstanding) == \
        (2, Decimal("131.00"), Decimal("91.00"), Decimal("40.00"))
    assert by_client["Example Firm LLP"].billed == Decimal("131.00")
    assert "Advocate & Partners" not in by_client
    assert list(by_month) == ["2026-01", "2026-02"] and by_month["2026-01"].paid == Decimal("91.00")
    assert len(ledger.invoices(status="paid")) == 1 and len(ledger.invoices(status="void")) == 1
    ledger.mark_unpaid("2026-0001")
    assert ledger.invoice("2026-0001").status == "open" and ledger.invoice("2026-0001").paid == 0


def test_activity_log(ledger):
    ledger.log_activity("agreement", case_name="Roe v. Doe", index_no="700001/2025", attorney="Alex", firm="Example Firm LLP",
                        file_path="a.pdf", ts="2026-03-01T10:00:00")
    ledger.log_activity("mofr", case_name="Roe v. Doe", ts="2026-03-02T11:00:00")
    ledger.log_activity("invoice", case_name="Smith v. Jones", pages=30, invoice_no="2026-0004", ts="2025-03-02T11:00:00")
    assert [a.kind for a in ledger.activity()] == ["mofr", "agreement", "invoice"]
    assert [a.kind for a in ledger.activity(kind="invoice")] == ["invoice"]
    assert len(ledger.activity(since="2026-01-01", until="2026-03-01")) == 1
    assert ledger.activity(text="smith")[0].pages == 30


def test_csv_mirror_and_exports(ledger, tmp_path):
    ledger.log_activity("mofr", case_name="Roe v. Doe")  # any change refreshes the CSV copies
    with open(tmp_path / "mirror" / "invoices.csv", encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    assert len(rows) == 4 and rows[0]["Invoice No."] == "2026-0003" and rows[0]["Date"] == "2/2/2026"
    assert (tmp_path / "mirror" / "activity.csv").read_text(encoding="utf-8-sig").count("MOFR") == 1

    from openpyxl import load_workbook
    wb = load_workbook(ledger.export_xlsx(tmp_path / "x.xlsx"))
    assert wb.sheetnames == ["Invoices", "Activity", "By firm", "By month"]
    assert wb["Invoices"].max_row == 5 and wb["By firm"]["A2"].value == "Example Firm LLP"

    html = ledger.export_html(tmp_path / "r.html", ledger.invoices(year=2026), "Invoices 2026").read_text("utf-8")
    assert "<svg" in html and "Invoices 2026" in html and "$203.00" in html and "Advocate &amp; Partners" in html
    assert "2025-0009" not in html


def test_locked_csv_is_not_fatal(ledger, tmp_path, monkeypatch):
    def locked(*a, **k):
        raise PermissionError("in use")
    monkeypatch.setattr(Ledger, "export_csv", locked)
    ledger.log_activity("mofr")  # logged, not raised
    assert ledger.activity()[0].kind == "mofr"
