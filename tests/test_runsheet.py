"""The run sheet: takes read from the reporters' initials on each transcript page, written to Excel, and added
to the case's run sheet when there is one (fictional names and cases)."""
from copy import copy
from datetime import date

import pymupdf
import pytest
from openpyxl import Workbook, load_workbook

from minute_filler.batch import fill_jobs, group, read_docs
from minute_filler.deliver import generate
from minute_filler.ingest import ingest_pdf
from minute_filler.runsheet import (FIRST, Row, RunSheetOpts, add_takes, choose, find_sheets, read_info,
                                    run_sheet_summary, same_name)
from minute_filler.takes import (PageMark, find_takes, initials_of, is_index_page, reporter_label, scan_page, scan_pdf,
                                 title_reporters, unspace)

from helpers import pat_settings

TITLE_1 = """ 1    SUPREME COURT OF THE STATE OF NEW YORK
      COUNTY OF QUEENS:  CIVIL TERM:  PART 14
 2    JANE ROE,                          Index No. {index}
 3                       Plaintiff,
 4             -against-
 5    {defendant},
 6                       Defendant.
 7    {day}
 8    B E F O R E :  HONORABLE MARIA T. ALVAREZ
 9    A P P E A R A N C E S :
10    ADVOCATE LAW GROUP
      Attorneys for the Plaintiff
11    BY: SAMUEL R. ADVOCATE, ESQ.
12    (Title continues on next page.)"""

TITLE_2 = """ 1    COUNSEL & WARD, LLP
      Attorneys for the Defendant
 2    BY: ALEX B. COUNSEL, ESQ.
 3
 4                     PAT REPORTER
 5                     DANA SMITH
                       Senior Court Reporters"""

SWORN = """ 1              THE CLERK:  Please state your name.
 2           J O H N   D O E, called as a witness by and on behalf
 3    of the Plaintiff, after having been first duly sworn,
 4    testified as follows:"""

EXCUSED = """ 1              THE COURT:  Thank you.  You may step down.
 2              (Whereupon, the witness stepped down.)"""

TALK = """ 1              THE COURT:  Good morning.
 2              MS. ADVOCATE:  Good morning, Your Honor.
 3              THE COURT:  Let's continue."""

# (page number, initials, running head, text): two title pages by Pat, three pages by Dana with a witness
# sworn in, three by Pat where he steps down, then Dana's take with the next witness - its last page
# lacks the initials - and the word index after the transcript.
DAY_1 = [
    (310, "pr", "", TITLE_1), (311, "pr", "", TITLE_2),
    (312, "ds", "Proceedings", TALK), (313, "ds", "Proceedings", SWORN), (314, "ds", "J. Doe - Plaintiff - Direct", TALK),
    (315, "pr", "J. Doe - Plaintiff - Direct", TALK), (316, "pr", "J. Doe - Plaintiff - Cross", EXCUSED),
    (317, "pr", "Proceedings", TALK),
    (318, "ds", "M. Lee - Defendant - Direct", TALK), (319, "", "M. Lee - Defendant - Direct", TALK),
    (None, "", "Roe v. Poe", "Min-U-Script\nabout (2) 312:4;315:9\nafter (1) 313:3"),
]


def transcript(path, pages=DAY_1, index="712345/2021", defendant="SAM POE", day="June 3, 2026"):
    """Writes a fictional transcript PDF: one page per (number, initials, head, text) in pages, the initials
    at the foot. index, defendant and day fill the title page. Returns path."""
    doc =pymupdf.open()
    for number, initials, head, text in pages:
        page = doc.new_page(width=612, height=792)
        if head:
            page.insert_text((150, 36), head, fontsize=10, fontname="cour")
        if number is not None:
            page.insert_text((300, 63), str(number), fontsize=10, fontname="cour")
        text = text.format(index=index, defendant=defendant, day=day)
        for i, line in enumerate(text.splitlines()):
            page.insert_text((40, 100 + 22 * i), line, fontsize=10, fontname="cour")
        if initials:
            page.insert_text((520, 740), initials, fontsize=10, fontname="cour")
    doc.save(path)
    return path


@pytest.fixture
def s(tmp_path):
    s = pat_settings()
    s.output_dir = str(tmp_path / "out")
    s.records_dir = str(tmp_path / "records")
    s.runsheet_dir = str(tmp_path / "sheets")
    return s


def job_of(path, s):
    """The one job made from a transcript."""
    docs, errors = read_docs([str(path)], s)
    assert not errors
    job, = group(docs, s)
    return job


def sheet_rows(path):
    """The rows of our run sheet's takes, columns A to K."""
    ws =load_workbook(path)["Run Sheet"]
    return [[ws.cell(r, c).value for c in range(1, 12)] for r in range(FIRST, ws.max_row + 1)]


# ------------------------------------------------------------------ reading the pages

def test_takes_follow_the_initials(tmp_path):
    with pymupdf.open(transcript(tmp_path / "t.pdf")) as doc:
        marks = scan_pdf(doc)
    assert len(marks) == 10  # the word index is not part of the transcript
    assert [m.number for m in marks] == list(range(310, 320)) and marks[0].initials == "pr" and not marks[9].initials
    assert marks[3].sworn == "JOHN DOE" and marks[6].excused and marks[4].witness() == "J. Doe"
    takes = find_takes(marks, title_count=2)
    assert [(t.initials, t.start, t.pages, t.title_only) for t in takes] == \
        [("pr", 310, 2, True), ("ds", 312, 3, False), ("pr", 315, 3, False), ("ds", 318, 2, False)]
    assert takes[1].witness_start == ["JOHN DOE"]       # sworn in: the head's "J. Doe" is the same witness
    assert takes[2].witness_start == [] and takes[2].witness_end == ["JOHN DOE"]
    assert takes[3].witness_start == ["M. Lee"]         # the next witness, by the running head


def test_names_for_the_initials():
    names = title_reporters("BY: ALEX B. COUNSEL, ESQ.\n\n 4   PAT REPORTER\n 5   DANA SMITH\n   Senior Court Reporters")
    assert names == ["PAT REPORTER", "DANA SMITH"]
    assert initials_of("Pat Q. Reporter") == {"pr", "pqr"}
    assert reporter_label("ds", names, {}, "Pat Reporter") == "Dana"
    assert reporter_label("pr", [], {}, "Pat Reporter") == "Pat"       # the user's own, from My info
    assert reporter_label("px", [], {}, "Pat Reporter", own_initials="PX") == "Pat"
    assert reporter_label("ds", names, {"ds": "Dana S."}, "Pat Reporter") == "Dana S."  # Settings win
    assert reporter_label("kl", names, {}, "Pat Reporter") == "KL"      # nobody known: the initials
    assert unspace("D R.  J A N E   R O E") == "DR. JANE ROE" and unspace("JOHN  DOE") == "JOHN DOE"


def test_the_word_index_is_not_billed(tmp_path, s):
    job = job_of(transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf"), s)
    assert job.docs[0].ing.page_count == 11 and job.transcript_pages() == 10
    job.page_basis[job.docs[0].key()] = "*"  # the whole transcript (Pat wrote 5 of its pages: see test_own_pages)
    assert job.invoice_pages() == 10
    note = job.docs[0].regex.fields["est_pages"][0].note
    assert "transcript pages 310-319" in note and "not counting 1 page(s) after the transcript (word index)" in note


# ------------------------------------------------------------------ the run sheet

def test_a_new_run_sheet(tmp_path, s):
    job = job_of(transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf"), s)
    opts = job.runsheet_opts(s)
    paths = generate(job.case, s, tmp_path / "out", ["runsheet"], runsheet=opts)
    assert paths == [opts.path] and opts.created and opts.added == 4
    assert opts.path.parent == tmp_path / "sheets" and opts.path.name == "June 2026 712345-2021 Jane Roe v. Sam Poe - Run Sheet.xlsx"
    rows = sheet_rows(opts.path)
    assert [(r[0].date(), r[2], r[5], r[8], r[9], r[10]) for r in rows] == [
        (date(2026, 6, 3), "Pat", 2, None, None, "title page only"),
        (date(2026, 6, 3), "Dana", 3, "John Doe", None, None),
        (date(2026, 6, 3), "Pat", 3, None, "John Doe", None),
        (date(2026, 6, 3), "Dana", 2, "M. Lee", None, None)]
    assert rows[0][6] == 310 and rows[1][6] == "=H5+1"  # the pages carry on from the row above
    assert rows[1][3].startswith('=IF(A6="","",IF(K6="title page only",0.5,COUNTIFS(')
    wb = load_workbook(opts.path)
    assert wb["DjinnIt"].sheet_state == "hidden" and [r.value for r in wb["By Reporter"]["A"]][:3] == \
        ["Reporter", "Pat", "Dana"]
    info = read_info(opts.path)
    assert info.ours and info.index_nos == ["712345/2021"] and info.case_name == "Jane Roe v. Sam Poe"
    assert "Started the run sheet" in run_sheet_summary(opts) and "4 takes" in run_sheet_summary(opts)


def test_the_next_day_is_added_in_order(tmp_path, s):
    day2 = [(320, "ds", "", TITLE_1), (321, "ds", "Proceedings", TALK), (322, "pr", "Proceedings", TALK)]
    second = job_of(transcript(tmp_path / "Transcript 6-4-2026 Roe v Poe.pdf", day2, day="June 4, 2026"), s)
    first = job_of(transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf"), s)
    a = second.runsheet_opts(s)
    add_takes(second.case, a, s, [])
    b = first.runsheet_opts(s)
    add_takes(first.case, b, s, [])  # "ask" with nobody to ask: the same index number is added to
    assert b.path == a.path and not b.created and b.added == 4
    rows = sheet_rows(a.path)
    assert [(r[0].day, r[2], r[5]) for r in rows] == [(3, "Pat", 2), (3, "Dana", 3), (3, "Pat", 3), (3, "Dana", 2),
                                                       (4, "Dana", 2), (4, "Pat", 1)]
    assert rows[4][6] == "=H8+1"  # day 2 starts at page 320, right after day 1's 319
    # the same transcript again: nothing is added twice
    again = first.runsheet_opts(s)
    add_takes(first.case, again, s, [])
    assert (again.added, again.skipped) == (0, 4) and len(sheet_rows(a.path)) == 6
    assert "already has these takes" in run_sheet_summary(again)


def test_another_index_number_of_the_same_trial(tmp_path, s):
    job = job_of(transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf"), s)
    add_takes(job.case, job.runsheet_opts(s), s, [])
    other = job_of(transcript(tmp_path / "Transcript 6-5-2026 Roe v Poe.pdf", index="700999/2022",
                              day="June 5, 2026"), s)
    found = find_sheets(other.case, [tmp_path / "sheets"])
    assert [f.why for f in found] == ["name"] and found[0].reason() == "same case name, index 712345/2021"
    assert choose(other.case, s, [tmp_path / "sheets"], None) is None  # "ask" with nobody to ask: a new one
    s.runsheet_existing = "add"
    assert choose(other.case, s, [tmp_path / "sheets"], None) == found[0].path
    opts = other.runsheet_opts(s)
    add_takes(other.case, opts, s, [])
    assert opts.path == found[0].path and read_info(opts.path).index_nos == ["712345/2021", "700999/2022"]
    unrelated = job_of(transcript(tmp_path / "Transcript 6-5-2026 Roe v Hill.pdf", index="700111/2022",
                                  defendant="TERRY HILL"), s)
    assert find_sheets(unrelated.case, [tmp_path / "sheets"]) == []
    assert same_name("Jane Roe, as Administrator of the Estate of Sam Roe v. Sam Poe, M.D., et al.",
                     "September 2026 712345-2021 ROE v Poe et al. - Page Sheet")


def test_adding_to_a_run_sheet_made_elsewhere(tmp_path, s):
    """A run sheet downloaded from Google Sheets: rows go at its end, its formulas are copied down."""
    path = tmp_path / "sheets" / "June 2026 712345-2021 ROE v Poe - Page Sheet.xlsx"
    path.parent.mkdir()
    wb = Workbook()
    ws = wb.active
    ws.append(["Write here", "Don't write here", "Write here", "", "", "Write here", "", "", "", "", ""])
    ws.append(["Date", "Weekday", "Reporter", "Day's Take", "Running Take", "Pages written", "Starting Page No.",
               "Ending Page no.", "Witness Start", "Witness End", "Note"])
    ws.append([date(2026, 6, 2), '=TEXT(A3,"dddd")', "Pat", 1, 1, 10, 300, "=G3+F3-1", None, None, "opening"])
    ws.append([None, '=IF(A4="","-",TEXT(A4,"dddd"))', None, None, None, None, "=H3+1", "=IF(F4=\"\",FALSE,G4+F4-1)"])
    wb.save(path)
    job = job_of(transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf"), s)
    found = find_sheets(job.case, [path.parent])
    assert [(f.path, f.ours, f.why) for f in found] == [(path, False, "index")]
    opts = job.runsheet_opts(s)
    add_takes(job.case, opts, s, [])
    assert opts.path == path and opts.added == 4
    ws = load_workbook(path).active
    assert [ws.cell(r, 3).value for r in range(3, 8)] == ["Pat", "Pat", "Dana", "Pat", "Dana"]
    assert ws["A4"].value.date() == date(2026, 6, 3) and ws["F4"].value == 2
    assert ws["G4"].value == "=H3+1" and ws["G5"].value == "=H4+1"  # their formula, copied down
    assert ws["B5"].value == '=IF(A5="","-",TEXT(A5,"dddd"))' and ws["K3"].value == "opening"
    assert (path.parent / "June 2026 712345-2021 ROE v Poe - Page Sheet (before DjinnIt).xlsx").exists()


def test_a_batch_puts_every_day_on_one_run_sheet(tmp_path, s):
    folder = tmp_path / "in"
    folder.mkdir()
    transcript(folder / "Transcript 6-3-2026 Roe v Poe.pdf")
    transcript(folder / "Transcript 6-4-2026 Roe v Poe.pdf", [(320, "ds", "", TITLE_1), (321, "pr", "", TALK)],
               day="June 4, 2026")
    (folder / "note.txt").write_text("Please send the minutes in Smith v Jones, Index No. 712222/2024, 5/22/2026.")
    docs, _ = read_docs([str(p) for p in sorted(folder.iterdir())], s)
    jobs = group(docs, s)
    assert len(jobs) == 3
    done = fill_jobs(jobs, s, outputs=["runsheet"])
    sheets = list((tmp_path / "sheets").glob("*.xlsx"))
    assert len(sheets) == 1 and done == sheets and len(sheet_rows(sheets[0])) == 6
    smith = next(j for j in jobs if "Smith" in j.title())
    assert smith.error == "no run sheet: no transcript PDF among the inputs" and not smith.saved


def test_a_scan_gets_one_row(tmp_path, s):
    job = job_of(transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf"), s)
    job.docs[0].ing.marks = []  # as for a scanned transcript: no text to find the initials in
    rows = job.runsheet_opts(s).rows
    assert len(rows) == 1 and rows[0].pages == 11 and rows[0].note == "reporter's initials not found"


def test_open_in_excel(tmp_path, s, monkeypatch):
    """A run sheet open in Excel can't be replaced: the message says to close it, and no half-saved copy stays."""
    job =job_of(transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf"), s)
    opts = job.runsheet_opts(s)
    add_takes(job.case, opts, s, [])
    import minute_filler.runsheet as rs

    def locked(src, dst):
        raise PermissionError(13, "The process cannot access the file")
    monkeypatch.setattr(rs.os, "replace", locked)
    again = RunSheetOpts([Row(date(2026, 6, 4), "Pat", 5, 320)], str(opts.path))
    with pytest.raises(PermissionError, match="close it in Excel"):
        add_takes(job.case, again, s, [])
    assert len(sheet_rows(opts.path)) == 4 and not list(opts.path.parent.glob("*.saving.xlsx"))


# ------------------------------------------------------------------ found in the run sheet sweep

D3, D4 = date(2026, 6, 3), date(2026, 6, 4)
NEXT_SWORN = """ 3           MARY LEE, called as a witness by and on behalf
 4    of the Defendant, having been first duly sworn,
 5    testified as follows:"""


def test_two_pages_saying_the_witness_stepped_down(tmp_path, s):
    """The second "stepped down" wrote a NUL character into the Witness End column: Excel refused the file."""
    pages = DAY_1[:6] + [(316, "pr", "J. Doe - Plaintiff - Cross", EXCUSED), (317, "pr", "Proceedings", EXCUSED)] \
        + DAY_1[8:]
    with pymupdf.open(transcript(tmp_path / "t.pdf", pages)) as doc:
        takes = find_takes(scan_pdf(doc), title_count=2)
    assert takes[2].witness_end == ["JOHN DOE"]
    job = job_of(tmp_path / "t.pdf", s)
    opts = job.runsheet_opts(s)
    add_takes(job.case, opts, s, [])
    assert [r[9] for r in sheet_rows(opts.path)] == [None, None, "John Doe", None]


def test_redirect_and_recross_heads_name_the_witness():
    assert PageMark(head="J. Doe - Plaintiff - Redirect").witness() == "J. Doe"
    assert PageMark(head="M. Lee - Defendant - Re-cross").witness() == "M. Lee"
    assert PageMark(head="Proceedings").witness() == ""


def test_one_witness_steps_down_and_the_next_is_sworn_on_the_same_page(tmp_path):
    pages = DAY_1[:5] + [(315, "pr", "J. Doe - Plaintiff - Cross", EXCUSED + "\n" + NEXT_SWORN),
                         (316, "ds", "M. Lee - Defendant - Direct", TALK)]
    with pymupdf.open(transcript(tmp_path / "t.pdf", pages)) as doc:
        marks = scan_pdf(doc)
        takes = find_takes(marks, title_count=2)
    assert marks[5].sworn == "MARY LEE" and marks[5].excused and not marks[5].sworn_first
    assert takes[2].witness_end == ["JOHN DOE"] and takes[2].witness_start == ["MARY LEE"]
    assert takes[3].witness_start == []  # still Mary Lee: not listed again
    # the other way round: sworn in and stepped down on one page
    marks = [PageMark(1, "pr"), PageMark(2, "pr", "Proceedings", sworn="MARY LEE", excused=True, sworn_first=True)]
    take, = find_takes(marks, title_count=1)
    assert take.witness_start == take.witness_end == ["MARY LEE"]


def test_a_witness_with_jr_after_the_name(tmp_path):
    sworn = SWORN.replace("J O H N   D O E,", "JOHN DOE, JR.,")
    pages = DAY_1[:3] + [(313, "ds", "Proceedings", sworn), (314, "ds", "J. Doe - Plaintiff - Direct", TALK)]
    with pymupdf.open(transcript(tmp_path / "t.pdf", pages)) as doc:
        marks = scan_pdf(doc)
    assert marks[3].sworn == "JOHN DOE, JR."
    takes = find_takes(marks, title_count=2)
    assert takes[1].witness_start == ["JOHN DOE, JR."]  # the head's "J. Doe" is the same witness


def test_reporters_on_the_title_page():
    # an attorneys' block right above, a name on the reporter's line, letters after a name
    assert title_reporters("ADVOCATE LAW GROUP\nAttorneys for the Plaintiff\nDANA SMITH\nSenior Court Reporter") \
        == ["DANA SMITH"]
    assert title_reporters("BY: ALEX B. COUNSEL, ESQ.\n\nPAT REPORTER, Senior Court Reporter") == ["PAT REPORTER"]
    assert title_reporters("COUNSEL & WARD, LLP\n\nDANA SMITH, RPR\nOfficial Court Reporter") == ["DANA SMITH"]
    # the block of names ends at a blank line
    assert title_reporters("JOHN DOE\n\nDANA SMITH\nSenior Court Reporter") == ["DANA SMITH"]
    assert title_reporters("PAT REPORTER\nSenior Court Reporter\n\nDANA SMITH\nSenior Court Reporter") == \
        ["PAT REPORTER", "DANA SMITH"]


def _numbered_page(doc, top: str | None, foot: str | None, text: str = "Q. And then?") -> None:
    """A page with its number at the top or the foot, and the line numbers 1 and 25 alone at the left edge."""
    page = doc.new_page(width=612, height=792)
    if top:
        page.insert_text((534, 30), top, fontsize=10, fontname="cour")
    page.insert_text((40, 80), "1", fontsize=10, fontname="cour")      # the first line's number, in the top margin
    page.insert_text((80, 300), text, fontsize=10, fontname="cour")
    page.insert_text((40, 720), "25", fontsize=10, fontname="cour")    # the last line's, in the bottom margin
    if foot:
        page.insert_text((300, 770), foot, fontsize=10, fontname="cour")
    page.insert_text((520, 745), "pr", fontsize=10, fontname="cour")


def test_page_number_at_the_foot(tmp_path):
    doc = pymupdf.open()
    _numbered_page(doc, None, "312")   # at the foot only: the "1" at the top is the first line's number
    _numbered_page(doc, None, "313")
    _numbered_page(doc, "3", None)     # at the top (a transcript starting at page 1): the "25" is a line number
    marks = [scan_page(p) for p in doc]
    assert [m.number for m in marks] == [312, 313, 3]


def test_the_word_index_is_found_by_its_references():
    condensed = "Min-U-Script\nJane Roe v. Sam Poe\nPage 5\nQ. Were you there at 9:00?\nA. Yes."
    assert not is_index_page(condensed)  # a condensed transcript prints Min-U-Script on every page
    assert not is_index_page("7:00 7:00 8:00 9:38 10:04 7:00 7:00 8:00 7:00 9:45")  # times, not references
    assert is_index_page("about (2) 312:4;315:9\nafter (3) 313:3;316:12;318:20\nagain (3) 314:5;317:8;319:2")
    assert is_index_page("Min-U-Script\nabout (2) 312:4;315:9\nafter (1) 313:3")


def test_a_condensed_transcript_is_counted_whole(tmp_path):
    pages = [(n, i, h, t + "\nMin-U-Script") for n, i, h, t in DAY_1[:-1]] + [DAY_1[-1]]
    with pymupdf.open(transcript(tmp_path / "t.pdf", pages)) as doc:
        assert len(scan_pdf(doc)) == 10


def _add_day_1(tmp_path, s):
    """Starts the run sheet with DAY_1's four takes; returns the job and the run sheet's path."""
    job =job_of(transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf"), s)
    opts = job.runsheet_opts(s)
    add_takes(job.case, opts, s, [])
    return job, opts.path


def test_the_starting_page_follows_the_row_with_pages(tmp_path, s):
    job, path = _add_day_1(tmp_path, s)
    wb = load_workbook(path)
    wb["Run Sheet"]["A9"], wb["Run Sheet"]["C9"] = D3, "Kim"  # a row typed in, without pages
    wb.save(path)
    opts = RunSheetOpts([Row(D4, "Dana", 2, 320)])
    add_takes(job.case, opts, s, [])
    rows = sheet_rows(path)
    assert [r[2] for r in rows] == ["Pat", "Dana", "Pat", "Dana", "Kim", "Dana"]
    assert rows[5][6] == "=H8+1"  # page 320 follows Dana's 319 on row 8, not the empty row 9


def test_what_was_typed_on_our_run_sheet_stays(tmp_path, s):
    job, path = _add_day_1(tmp_path, s)
    wb = load_workbook(path)
    ws = wb["Run Sheet"]
    ws["F6"] = "=1+2"                                # Dana's 3 pages, typed as a sum
    ws["A9"], ws["C9"], ws["F9"] = "TBD", "Kim", 4   # a row with no date yet
    ws["L6"] = "=F6*2"                               # the user's own column
    wb.save(path)
    opts = RunSheetOpts([Row(date(2026, 6, 2), "Dana", 10, 300)])  # the day before: goes first, the rest move down
    add_takes(job.case, opts, s, [])
    ws = load_workbook(path)["Run Sheet"]
    assert [ws.cell(r, 3).value for r in range(5, 11)] == ["Dana", "Pat", "Dana", "Pat", "Dana", "Kim"]
    assert ws["F7"].value == "=1+2" and ws["L7"].value == "=F7*2"  # moved with their row
    assert ws["A10"].value == "TBD" and ws["F10"].value == 4         # the undated row stays where it was


def test_merged_cells_and_notes_beside_the_title(tmp_path, s):
    job, path = _add_day_1(tmp_path, s)
    wb = load_workbook(path)
    ws = wb["Run Sheet"]
    ws.merge_cells("A1:K1")
    ws["M1"] = "Ask Dana about 6/5"
    ws.merge_cells("M1:O1")
    wb.save(path)
    add_takes(job.case, RunSheetOpts([Row(D4, "Dana", 2, 320)]), s, [])
    ws = load_workbook(path)["Run Sheet"]
    assert ws["M1"].value == "Ask Dana about 6/5" and len(sheet_rows(path)) == 5


def test_by_reporter_with_the_tab_renamed_and_names_in_capitals(tmp_path, s):
    job, path = _add_day_1(tmp_path, s)
    wb = load_workbook(path)
    wb["Run Sheet"].title = "Roe's trial"
    wb.save(path)
    add_takes(job.case, RunSheetOpts([Row(D4, "DANA", 2, 320)]), s, [])
    by = load_workbook(path)["By Reporter"]
    assert [by.cell(r, 1).value for r in range(2, 4)] == ["Pat", "Dana"] and by["A4"].value == "Total"
    assert by["C3"].value.startswith("=SUMIF('Roe''s trial'!$C$5:$C$5000,A3,")


def test_nothing_new_leaves_the_run_sheet_alone(tmp_path, s, monkeypatch):
    job, path = _add_day_1(tmp_path, s)
    import minute_filler.runsheet as rs

    def locked(src, dst):
        raise PermissionError(13, "The process cannot access the file")
    monkeypatch.setattr(rs.os, "replace", locked)  # open in Excel: no matter, there is nothing to write
    again = job.runsheet_opts(s)
    add_takes(job.case, again, s, [])
    assert (again.added, again.skipped) == (0, 4)


def test_the_same_transcript_twice_in_one_go(tmp_path, s):
    job = job_of(transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf"), s)
    rows = job.runsheet_opts(s).rows
    opts = RunSheetOpts(rows + [copy(r) for r in rows])
    add_takes(job.case, opts, s, [])
    assert (opts.added, opts.skipped) == (4, 4) and len(sheet_rows(opts.path)) == 4


def test_a_locked_file_is_tried_again(tmp_path, s, monkeypatch):
    import minute_filler.runsheet as rs
    tries = []
    real = rs.os.replace

    def busy(src, dst):  # Dropbox keeping the file to itself for a moment
        tries.append(1)
        if len(tries) < 3:
            raise PermissionError(13, "The process cannot access the file")
        real(src, dst)
    monkeypatch.setattr(rs.os, "replace", busy)
    monkeypatch.setattr(rs.time, "sleep", lambda t: None)
    job, path = _add_day_1(tmp_path, s)
    assert len(tries) == 3 and len(sheet_rows(path)) == 4


def test_case_names_on_run_sheets():
    assert same_name("Jane Roe v. Sam Poe", "June 2026 712345-2021 Jane Roe v. Sam Poe - Run Sheet")
    assert same_name("Roe v. The City of New York", "June 2026 712345-2021 Roe v. The City of New York - Run Sheet")
    assert same_name("Jane Roe v. Acme Widget Corporation", "Roe v Acme Widget Corp. - Page Sheet")
    assert not same_name("Maria Garcia v. NYC Transit Authority", "Maria Rodriguez v. NYC Transit Authority")
    assert not same_name("Jane Roe v. Sam Poe", "Jane Smith v. Sam Jones")


def _their_sheet(path, headings, rows):
    """A run sheet made elsewhere (not by this app): a "Write here" row, the headings, then rows."""
    path.parent.mkdir(exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.append(["Write here"] * len(headings))
    ws.append(headings)
    for r in rows:
        ws.append(r)
    wb.save(path)
    return path


HEADINGS = ["Date", "Weekday", "Reporter", "Day's Take", "Running Take", "Pages written", "Starting Page No.",
            "Ending Page no.", "Witness Start", "Witness End", "Note"]


def test_days_typed_in_ahead_on_a_run_sheet_made_elsewhere(tmp_path, s):
    path = _their_sheet(tmp_path / "sheets" / "June 2026 712345-2021 ROE v Poe - Page Sheet.xlsx", HEADINGS, [
        [date(2026, 6, 2), None, "Pat", 1, 1, 10, 300, "=G3+F3-1"],
        [D3], [D4]])  # the next days, typed in ahead
    job = job_of(transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf"), s)
    opts = job.runsheet_opts(s)
    add_takes(job.case, opts, s, [])
    ws = load_workbook(path).active
    assert [ws.cell(r, 3).value for r in range(3, 9)] == ["Pat", "Pat", None, "Dana", "Pat", "Dana"]
    assert ws["A4"].value.date() == D3 and ws["A5"].value.date() == D4  # June 4 is still there, for its day
    assert ws["G6"].value == 312  # not "=H5+1": row 5 has no pages yet


def test_headings_of_a_run_sheet_made_elsewhere(tmp_path, s):
    path = _their_sheet(tmp_path / "sheets" / "June 2026 712345-2021 ROE v Poe - Page Sheet.xlsx",
                        ["Trial Date", "Reporter", "No. of\nPages", "Starting Page", "Ending Page"],
                        [[date(2026, 6, 2), "Pat", 10, 300, 309]])
    job = job_of(transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf"), s)
    add_takes(job.case, job.runsheet_opts(s), s, [])
    ws = load_workbook(path).active
    assert [(ws.cell(r, 1).value.date(), ws.cell(r, 2).value, ws.cell(r, 3).value) for r in (4, 5)] == \
        [(D3, "Pat", 2), (D3, "Dana", 3)]
    bad = _their_sheet(tmp_path / "other" / "June 2026 712345-2021 ROE v Poe - Page Sheet.xlsx",
                       ["Date", "Reporter", "Starting Page", "Ending Page"], [])
    with pytest.raises(ValueError, match="no Pages column"):
        add_takes(job.case, RunSheetOpts(job.runsheet_opts(s).rows, str(bad)), s, [])


def test_what_is_not_a_run_sheet(tmp_path):
    from minute_filler.invoice import TEMPLATE
    assert read_info(TEMPLATE) is None  # the invoice spreadsheet's Page Log
    path = _their_sheet(tmp_path / "June 2026 712345-2021 ROE v Poe - Page Sheet.xlsx", HEADINGS,
                        [[date(2026, 6, 2), None, "Pat", 1, 1, 10, 300]])
    info = read_info(path)
    assert info.index_nos == ["712345/2021"]  # the date 2026-06-02 is not index number 2026/2006
