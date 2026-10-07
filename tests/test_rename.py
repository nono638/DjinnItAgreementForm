"""The app was DjinnItAgreementForm before 2.0: the first start under the new name carries the old settings
folder over (settings, rate sheets, records, signature) and renames the folders under Documents named after it,
and what the old version made (its PDFs, run sheets and settings files) is still recognised. Also the yin-yang
that took the djinn's place: it swirls while the app works and while it waits for documents, goes on a while
after a quick read before the "ready" picture, and swirls once when the still picture is clicked. All names are
made up; every folder is a temporary one (conftest.isolated_appdata)."""
import json
import os
from pathlib import Path

import pymupdf
from openpyxl import Workbook

from minute_filler import settings as settings_mod
from minute_filler.fill import is_generated
from minute_filler.runsheet import FORMAT, META, _meta, meta_sheet, read_info
from minute_filler.settings import Settings, settings_dir


def old_install() -> tuple[Path, Path]:
    """A DjinnIt user's folders: %APPDATA%\\DjinnItAgreementForm with settings, a rate sheet, the records and a
    signature, and Documents\\DjinnIt Records and Documents\\DjinnIt Run Sheets with a file each."""
    appdata = Path(os.environ["APPDATA"]) / "DjinnItAgreementForm"
    (appdata / "Rate Sheets").mkdir(parents=True)
    (appdata / "Rate Sheets" / "Mine.csv").write_text("speed,original\nRegular,4.30\n", encoding="utf-8")
    (appdata / "records.db").write_bytes(b"records")
    (appdata / "signature.png").write_bytes(b"png")
    (appdata / "settings.json").write_text(json.dumps({
        "settings_version": 9, "profile": {"name": "Pat Reporter"}, "show_djinn": False,
        "signature_image": str(appdata / "signature.png"), "sign_reporter": True}), encoding="utf-8")
    docs = Path.home() / "Documents"
    for name in ("DjinnIt Records", "DjinnIt Run Sheets"):
        (docs / name).mkdir(parents=True)
        (docs / name / "kept.txt").write_text("mine", encoding="utf-8")
    return appdata, docs


def test_the_old_settings_folder_is_carried_over_and_the_documents_folders_renamed():
    old, docs = old_install()
    new = settings_dir()
    assert new.name == "YinItAgreementForm" and new.parent == old.parent
    assert (new / "Rate Sheets" / "Mine.csv").exists() and (new / "records.db").read_bytes() == b"records"
    assert (old / "settings.json").exists()  # the old folder is left as it was
    s = Settings.load()
    assert s.profile.name == "Pat Reporter" and s.show_yin is False  # (was show_djinn)
    assert Path(s.signature_image) == new / "signature.png" and s.signature()  # follows the folder
    assert (docs / "YinIt Records" / "kept.txt").exists() and not (docs / "DjinnIt Records").exists()
    assert (docs / "YinIt Run Sheets" / "kept.txt").exists()
    assert s.records_folder() == docs / "YinIt Records" and s.folder_for("runsheet") == docs / "YinIt Run Sheets"
    s.profile.name = "Changed"
    s.save()
    assert Settings.load().profile.name == "Changed"  # carried over once: not copied again


def test_a_documents_folder_that_cant_be_renamed_is_kept_where_it_is(monkeypatch):
    old, docs = old_install()

    def busy(self, target):
        raise PermissionError("open in Explorer")
    monkeypatch.setattr(Path, "rename", busy)
    settings_dir()
    s = Settings.load()
    assert s.records_folder() == docs / "DjinnIt Records"
    assert s.folder_for("runsheet") == docs / "DjinnIt Run Sheets"  # its run sheets are still found


def test_the_records_follow_the_renamed_folders():
    """The records kept pointing at Documents\\DjinnIt Run Sheets (and the old settings folder) after the
    rename: Open file and Print in Records found nothing there."""
    from minute_filler.records import Invoice, Ledger
    old, docs = old_install()
    (old / "records.db").unlink()
    lg = Ledger(old / "records.db")
    sheet = str(docs / "djinnit run sheets" / "Roe run sheet.xlsx")  # (Windows paths: any case)
    lg.log_activity("runsheet", case_name="Jane Roe v. Sam Poe", file_path=sheet,
                    origin=json.dumps({"sources": [str(docs / "DjinnIt Records" / "Roe.pdf")]}))
    elsewhere = str(docs / "Elsewhere" / "a.pdf")  # (not in a folder that was renamed: left alone)
    lg.log_activity("agreement", case_name="Jane Roe v. Sam Poe", file_path=elsewhere)
    lg.add_invoice(Invoice("2026-0001", "2026-06-03", amounts={"Regular": "63.00"},
                           file_path=str(old / "Invoices" / "Invoice 2026-0001.pdf")))
    new = settings_dir()
    got = Ledger(new / "records.db")
    paths = {a.kind: a.file_path for a in got.activity()}
    assert paths == {"runsheet": str(docs / "YinIt Run Sheets" / "Roe run sheet.xlsx"), "agreement": elsewhere}
    runsheet, = [a for a in got.activity() if a.kind == "runsheet"]
    assert json.loads(runsheet.origin)["sources"] == [str(docs / "YinIt Records" / "Roe.pdf")]
    assert got.invoice("2026-0001").file_path == str(new / "Invoices" / "Invoice 2026-0001.pdf")
    assert sheet in {a.file_path for a in Ledger(old / "records.db").activity()}  # the old copy is left as it was


def test_a_carry_over_that_couldnt_copy_the_settings_is_tried_again(monkeypatch):
    """A file that couldn't be copied (open in another program) stopped the carry-over half way, and the next
    start found the half-made folder and never carried over again."""
    import shutil
    old, docs = old_install()
    real = shutil.copytree
    locked = {"settings.json"}

    def copytree(src, dst, *a, **k):
        def copy(s, d, **kw):
            if Path(s).name in locked:
                raise PermissionError(13, "in use", s)
            return shutil.copy2(s, d, **kw)
        if len(a) >= 3:  # (copytree calls itself for each folder, with every argument in a row)
            return real(src, dst, *a[:2], copy, *a[3:], **k)
        return real(src, dst, *a, **{**k, "copy_function": copy})
    monkeypatch.setattr(shutil, "copytree", copytree)
    new = settings_dir()
    assert not (new / "settings.json").exists() and not (new / "records.db").exists()
    assert not list(new.parent.glob("YinItAgreementForm.partial*"))  # nothing half made is left
    assert (docs / "DjinnIt Records").exists()  # nothing renamed yet either
    settings_mod._CARRY_TRIED.clear()  # the next start
    locked = {"signature.png"}  # a file that isn't needed: the rest is carried over without it
    new = settings_dir()
    assert Settings.load().profile.name == "Pat Reporter" and (new / "Rate Sheets" / "Mine.csv").exists()
    assert not (new / "signature.png").exists() and (docs / "YinIt Records").exists()


def test_a_settings_file_exported_by_djinnit_names_the_renamed_folders(tmp_path):
    """Folders under Documents\\DjinnIt Records / Run Sheets in an exported file are read as the YinIt ones (a
    run sheet folder that isn't there may be left blank); the old names are never kept."""
    docs = Path.home() / "Documents"
    (docs / "YinIt Records").mkdir(parents=True)
    exported = tmp_path / "DjinnIt settings.json"
    exported.write_text(json.dumps({"app": "DjinnItAgreementForm", "personal": True, "settings_version": 9,
                                    "records_dir": str(docs / "DjinnIt Records"),
                                    "output_dirs": {"runsheet": str(docs / "DjinnIt Run Sheets" / "2026")}}),
                        encoding="utf-8")
    s = Settings.import_from(exported, Settings())
    assert s.records_dir == str(docs / "YinIt Records")
    assert s.output_dirs.get("runsheet", "") in ("", str(docs / "YinIt Run Sheets" / "2026"))
    assert "DjinnIt" not in json.dumps(s.output_dirs)


def test_a_new_user_gets_the_new_folders():
    s = Settings.load()
    assert settings_dir().name == "YinItAgreementForm" and s.show_yin is True
    assert s.records_folder().name == "YinIt Records" and s.folder_for("runsheet").name == "YinIt Run Sheets"
    assert not (Path(os.environ["APPDATA"]) / "DjinnItAgreementForm").exists()


def test_a_settings_file_exported_by_djinnit_can_be_imported(tmp_path):
    exported = tmp_path / "DjinnIt settings.json"
    exported.write_text(json.dumps({"app": "DjinnItAgreementForm", "personal": True, "settings_version": 9,
                                    "profile": {"name": "Dana Smith"}, "show_djinn": False}), encoding="utf-8")
    s = Settings.import_from(exported, Settings())
    assert s.profile.name == "Dana Smith" and s.show_yin is False


def test_pdfs_made_by_djinnit_are_still_known_as_the_apps_own(tmp_path):
    from minute_filler.gui.preview import file_kind
    path = tmp_path / "Minute Agreement - Roe v Poe.pdf"
    doc = pymupdf.open()
    doc.new_page()
    doc.set_metadata({"creator": "DjinnIt agreement"})
    doc.save(path)
    doc.close()
    assert is_generated(path) and file_kind(path) == "agreement"  # (not read again as an input)


def test_a_run_sheet_made_by_djinnit_is_found_and_takes_the_new_name(tmp_path):
    path = tmp_path / "Roe v Poe run sheet.xlsx"
    wb = Workbook()
    wb.active.title = "Run Sheet"
    old = wb.create_sheet("DjinnIt")
    for row in (["format", "DjinnIt run sheet 1"], ["case", "Jane Roe v. Sam Poe"], ["index", "123456/2025"]):
        old.append(row)
    wb.save(path)
    found = read_info(path)
    assert found and found.case_name == "Jane Roe v. Sam Poe" and found.index_nos == ["123456/2025"]
    _meta(wb, "Jane Roe v. Sam Poe", ["123456/2025"], {})  # written again: the hidden sheet is the new one
    assert meta_sheet(wb) == META and "DjinnIt" not in wb.sheetnames and wb[META]["B1"].value == FORMAT


def test_the_old_names_are_kept_only_where_old_files_need_them():
    """The app's name is the new one; the old one is kept in OLD_APP_NAMES, for the carry-over."""
    assert settings_mod.APP_NAME == "YinItAgreementForm"
    assert settings_mod.OLD_APP_NAMES[0] == "DjinnItAgreementForm"


def test_the_yin_yang_swirls_while_working_and_stops_when_done(make_window, qt):
    from PySide6.QtGui import QMovie
    from helpers import pat_settings
    s = pat_settings()
    s.use_ai, s.welcomed = False, True
    drop = make_window(s).drop
    assert drop.movie.isValid() and drop.movie.frameCount() > 1
    drop.set_mood("working")
    assert drop.movie.state() == QMovie.Running
    drop.set_mood("done")
    assert drop.movie.state() == QMovie.NotRunning and not drop.yin.pixmap().isNull()
    drop.set_mood(None)
    assert drop.movie.state() == QMovie.NotRunning
    drop.set_mood("working")
    drop.set_mood(None)
    drop.set_mood("working")  # hidden and shown again: it swirled no more (the old mood was kept as shown)
    assert drop.movie.state() == QMovie.Running


def test_the_yin_yang_swirls_while_the_app_waits(make_window, qt):
    """It swirls while waiting for documents too (it is the user's favourite), without saying "Working on it…"."""
    from PySide6.QtGui import QMovie
    from helpers import pat_settings
    s = pat_settings()
    s.use_ai, s.welcomed, s.show_yin = False, True, True
    win = make_window(s)
    win._set_status("Drop a document to begin", "")
    assert win.drop.movie.state() == QMovie.Running and win.drop.mood[0] == "idle"
    assert not win.drop.yin.pixmap().isNull() and "Working" not in win.drop.yin.toolTip()
    win._set_status("Reading…", "busy")
    assert win.drop.movie.state() == QMovie.Running and win.drop.mood[0] == "working"


def test_a_quick_read_swirls_a_while_before_the_ready_picture(make_window, qt):
    """A document read in a blink showed the swirl for a blink: it goes on for linger_ms; a problem shows at
    once, and so does anything asked for meanwhile."""
    import time
    from PySide6.QtGui import QMovie
    from PySide6.QtWidgets import QApplication
    from helpers import pat_settings
    s = pat_settings()
    s.use_ai, s.welcomed = False, True
    drop = make_window(s).drop
    drop.linger_ms = 300
    drop.set_mood("working")
    drop.set_mood("done")
    assert drop.mood[0] == "working" and drop.movie.state() == QMovie.Running  # still swirling
    end = time.monotonic() + 3
    while drop.mood[0] != "done" and time.monotonic() < end:
        QApplication.processEvents()
        time.sleep(0.01)
    assert drop.mood[0] == "done" and drop.movie.state() == QMovie.NotRunning
    drop.set_mood("working")
    drop.set_mood("stumped")  # a problem: at once
    assert drop.mood[0] == "stumped"
    drop.set_mood("working")
    drop.set_mood("done")
    drop.set_mood("working")  # another document before the swirl ended: no "ready" picture in between
    QApplication.processEvents()
    time.sleep(0.4)
    QApplication.processEvents()
    assert drop.mood[0] == "working"


def test_a_click_on_the_still_yin_yang_swirls_it_for_a_turn(make_window, qt):
    """Just for fun: a click on the ready (or stumped) picture swirls it once, then that picture is back."""
    import time
    from PySide6.QtCore import QPoint, Qt
    from PySide6.QtGui import QMovie
    from PySide6.QtTest import QTest
    from PySide6.QtWidgets import QApplication
    from helpers import pat_settings
    s = pat_settings()
    s.use_ai, s.welcomed = False, True
    drop = make_window(s).drop
    drop.linger_ms = 300
    drop.set_mood("stumped")
    QTest.mouseClick(drop.yin, Qt.LeftButton, pos=drop.yin.rect().center())
    assert drop.mood[0] == "working" and drop.movie.state() == QMovie.Running and drop.wanted == "stumped"
    end = time.monotonic() + 3
    while drop.mood[0] != "stumped" and time.monotonic() < end:
        QApplication.processEvents()
        time.sleep(0.01)
    assert drop.mood[0] == "stumped" and drop.movie.state() == QMovie.NotRunning
    QTest.mouseClick(drop.yin, Qt.LeftButton, pos=QPoint(1, 1))  # beside the picture: nothing
    assert drop.mood[0] == "stumped"
    drop.set_mood("idle")  # already swirling: a click changes nothing
    QTest.mouseClick(drop.yin, Qt.LeftButton, pos=drop.yin.rect().center())
    assert drop.mood[0] == "idle" and drop.wanted == "idle"
