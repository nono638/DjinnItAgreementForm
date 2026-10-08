"""Bugs found by the business-logic sweep, each pinned down: a page of chart times taken for the word index, the
run sheet crediting the unsigned first pages to the wrong reporter, an orphaned invoice keeping its number, a
run sheet locked for a moment being started again, one date for the whole invoice, two empty jobs told apart,
the dates of a trial as a range on the forms (and on a second line when they still don't fit), and the agreement
date filled in as a default. Then what the review of those fixes found: a range cut off at its last date, dates
split inside a date, a half-written invoice giving its number to the next, a detailed copy dated a day late, a
workbook that isn't a run sheet stopping the case's, case names in the log, and a word index taken for testimony.
All names made up."""
import re
import sqlite3
import zipfile
from datetime import date
from pathlib import Path

import openpyxl
import pymupdf
import pytest

from minute_filler import deliver, invoice, log, records
from minute_filler.batch import Job, group, read_docs
from minute_filler.fill import fill, form_path
from minute_filler.forms import original_map, ucs_map
from minute_filler.invoice import InvoiceOpts, make_invoice
from minute_filler.models import SRC_DEFAULT, Attorney
from minute_filler.mofr import fill_mofr
from minute_filler.records import Ledger
from minute_filler.runsheet import NOT_FOUND, Row, RunSheetOpts, SheetUnreadable, add_takes, find_sheets, read_info
from minute_filler.takes import has_line_numbers, is_index_page, scan_pdf

from helpers import ROE, make_case, pat_settings
from test_runsheet import DAY_1, TALK, TITLE_1, TITLE_2, transcript

ALEX = Attorney(name="Alex B. Counsel", firm="Counsel & Counsel", checked=True)
PR, DS = "pr", "ds"
# A page of testimony reading out the times on a chart: eight "h:mm" with minutes that could be line numbers
CHART = "\n".join(f"{i:2}    {line}" for i, line in enumerate([
    "Q.  What does the chart say about that morning?",
    "A.  Vitals: 9:10, 9:15, 9:20 and 9:25.",
    "Q.  And after that?",
    "A.  Again: 10:10, 10:15, 10:20 and 11:15.",
    "Q.  Nothing in between?",
    "A.  Not that I see.",
    "(Pause.)",
    "Q.  Go on."] + ["A.  That's all the chart says."] * 17, 1))


@pytest.fixture
def s(tmp_path):
    s = pat_settings()
    s.output_dir = str(tmp_path / "out")
    s.records_dir = str(tmp_path / "records")
    s.runsheet_dir = str(tmp_path / "sheets")
    return s


def job_of(path, s):
    """The one job made from a transcript."""
    docs, errors = read_docs([str(path)], s)
    assert not errors
    job, = group(docs, s)
    return job


# ------------------------------------------------------------------ pages and takes

def test_a_page_of_chart_times_is_not_the_word_index(tmp_path, s):
    """Eight times on one page look like eight page:line references; the page's line numbers (and the reporter's
    initials at its foot) say it is testimony, so the pages after it are still counted and billed."""
    assert not is_index_page(CHART)
    assert not is_index_page("9:10 a.m., 9:15 a.m., 9:20 a.m., 9:25 p.m., at 10:10, about 10:15, by 10:20, at 11:15")
    assert is_index_page("about (2) 312:4;315:9\nafter (3) 313:3;316:12;318:20\nagain (3) 314:5;317:8;319:2")
    pages = [(310, PR, "", TITLE_1), (311, PR, "Proceedings", CHART), (312, PR, "Proceedings", TALK),
             (313, PR, "Proceedings", TALK), DAY_1[-1]]  # the word index after the transcript still ends it
    path = transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf", pages)
    with pymupdf.open(path) as doc:
        assert [m.number for m in scan_pdf(doc)] == [310, 311, 312, 313]
    job = job_of(path, s)
    assert job.transcript_pages() == 4 and job.invoice_pages() == 4


def test_the_run_sheet_credits_the_unsigned_first_pages_as_whose_pages_said(tmp_path, s):
    """Three pages before the first initials, then Dana's and Pat's: Whose pages... says the first are Pat's, and
    the run sheet's takes agree with the invoice's count; "none" keeps them as a take of their own."""
    pages = [(310, "", "", TITLE_1), (311, "", "", TITLE_2), (312, "", "Proceedings", TALK)]
    pages += [(313 + i, DS, "Proceedings", TALK) for i in range(57)]
    pages += [(370 + i, PR, "Proceedings", TALK) for i in range(60)]
    job = job_of(transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf", pages), s)
    key = job.docs[0].key()
    job.front_owner[key] = PR
    rows = job.runsheet_opts(s).rows
    assert [(r.initials, r.reporter, r.pages) for r in rows] == [(PR, "Pat", 3), (DS, "Dana", 57), (PR, "Pat", 60)]
    counted = {}
    for r in rows:
        counted[r.initials] = counted.get(r.initials, 0) + r.pages
    assert counted == job.reporter_pages() == {PR: 63, DS: 57}
    job.front_owner[key] = "none"
    rows = job.runsheet_opts(s).rows
    assert [(r.initials, r.pages, r.note) for r in rows] == [("", 3, NOT_FOUND), (DS, 57, ""), (PR, 60, "")]
    assert job.reporter_pages() == {"": 3, DS: 57, PR: 60}


# ------------------------------------------------------------------ invoices and their numbers

def test_a_number_stays_taken_while_its_orphan_invoice_is_on_disk(s, tmp_path, monkeypatch):
    """The records refuse the invoice and its PDF can't be deleted: the number stays taken, so the next invoice
    doesn't share it with the file left behind."""
    case = make_case({**ROE, "est_pages": "30"}, attorneys=[ALEX])
    ledger = Ledger(tmp_path / "r.db")

    def refused(self, inv, year=0, seq=0):
        raise sqlite3.OperationalError("database is locked")

    def held(self, missing_ok=False):
        raise PermissionError(13, f"The process cannot access the file {self}")
    monkeypatch.setattr(Ledger, "add_invoice", refused)
    monkeypatch.setattr(Path, "unlink", held)
    with pytest.raises(sqlite3.OperationalError):
        make_invoice(case, ALEX, s, tmp_path / "out", InvoiceOpts(30), ledger)
    year = date.today().year
    orphan, = (tmp_path / "out").glob("*.pdf")
    assert f"{year}-0001" in orphan.name
    assert ledger.next_invoice_no()[0] == f"{year}-0002"  # (before: 0001 again, and two invoices with it)


def test_a_file_half_written_is_deleted_and_the_number_given_back(s, tmp_path, monkeypatch):
    """Saving the PDF fails part way (the disk full): nothing is left on disk, so the number is free again."""
    case = make_case({**ROE, "est_pages": "30"}, attorneys=[ALEX])
    ledger = Ledger(tmp_path / "r.db")

    def half(self, path, *args, **kwargs):
        Path(path).write_bytes(b"%PDF-1.7 half of it")
        raise OSError(28, "No space left on device")
    monkeypatch.setattr(pymupdf.Document, "save", half)
    with pytest.raises(OSError):
        make_invoice(case, ALEX, s, tmp_path / "out", InvoiceOpts(30), ledger)
    assert not list((tmp_path / "out").glob("*.pdf"))
    assert ledger.next_invoice_no()[0] == f"{date.today().year}-0001"


def test_an_invoice_made_across_midnight_has_one_date(s, tmp_path, monkeypatch):
    """The number's year, the date on the page and the records' date are one day, asked once: on December 31
    at midnight they would otherwise be a 2026 number on an invoice dated 2027."""
    class Midnight(date):
        calls = 0

        @classmethod
        def today(cls):
            cls.calls += 1
            return date(2026, 12, 31) if cls.calls == 1 else date(2027, 1, 1)
    monkeypatch.setattr(invoice, "date", Midnight)
    monkeypatch.setattr(records, "date", Midnight)
    case = make_case({**ROE, "est_pages": "30"}, attorneys=[ALEX])
    ledger = Ledger(tmp_path / "r.db")
    path, number = make_invoice(case, ALEX, s, tmp_path / "out", InvoiceOpts(30), ledger)
    inv, = ledger.invoices()
    assert number == "2026-0001" and inv.created == "2026-12-31"
    with pymupdf.open(path) as doc:
        dated = {w.field_name: w.field_value for w in doc[0].widgets()}["date"]
    assert dated == "12/31/2026"


# ------------------------------------------------------------------ run sheets

def test_a_run_sheet_locked_for_a_moment_is_not_started_again(tmp_path, s, monkeypatch):
    """The case's run sheet can't be opened (a sync holds it): the takes are not put on a second run sheet, the
    error names the file, and nothing new is made."""
    job = job_of(transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf"), s)
    first = job.runsheet_opts(s)
    add_takes(job.case, first, s, [])
    real = openpyxl.load_workbook

    def locked(path, *args, **kwargs):
        if Path(path).name == first.path.name:
            raise PermissionError(13, "The process cannot access the file")
        return real(path, *args, **kwargs)
    monkeypatch.setattr(openpyxl, "load_workbook", locked)
    again = RunSheetOpts([Row(date(2026, 6, 4), "Pat", 5, 320)])
    with pytest.raises(OSError, match="Run Sheet.xlsx could not be read"):  # (runsheet.SheetUnreadable)
        add_takes(job.case, again, s, [])
    assert not again.created and again.path is None
    assert [p.name for p in (tmp_path / "sheets").glob("*.xlsx")] == [first.path.name]  # no "(2)"


# ------------------------------------------------------------------ jobs

def test_two_empty_jobs_are_told_apart():
    """New job twice: the list finds each by itself, not by its (equal) contents."""
    a, b = Job(), Job()
    assert a != b and a == a and [a, b].index(b) == 1
    jobs = [a, b]
    jobs.remove(b)
    assert jobs == [a] and jobs[0] is a


# ------------------------------------------------------------------ the dates on the forms

FIVE = "9/28/2026, 9/29/2026, 9/30/2026, 10/1/2026, 10/2/2026"
GAPS = "9/28/2026, 9/30/2026, 10/2/2026, 10/5/2026, 10/7/2026"


def test_days_in_a_row_become_a_range():
    """Three or more days in a row are one range (an en dash); a gap keeps the list; two days stay as they are,
    and so does a text that is more than dates."""
    from minute_filler.fill import date_ranges
    assert date_ranges(FIVE) == "9/28/2026–10/2/2026"
    assert date_ranges("9/28/2026, 9/29/2026, 9/30/2026 and 10/5/2026") == "9/28/2026–9/30/2026, 10/5/2026"
    assert date_ranges("10/5/2026, 9/29/2026, 9/28/2026, 9/30/2026") == "9/28/2026–9/30/2026, 10/5/2026"
    assert date_ranges("9/28/2026, 9/29/2026") == "9/28/2026, 9/29/2026"
    assert date_ranges(GAPS) == GAPS
    assert date_ranges("6/3/2026 (a.m. session), 6/4/2026, 6/5/2026") == "6/3/2026 (a.m. session), 6/4/2026, 6/5/2026"
    assert date_ranges("") == ""


def _dates_fields(path, key_of):
    """(the dates field, the second dates line) of a filled form, by our keys."""
    with pymupdf.open(path) as doc:
        by_key = {key_of(w.field_name): (w.field_value, w.text_fontsize, w.rect) for w in doc[0].widgets()}
    return by_key["dates"], by_key["dates_2"]


def _width(text: str, fs: float) -> float:
    """The width a text takes on a form, measured with the fields' font itself, not with fill.text_width (which
    once measured "9/28/2026–10/2/2026" 14 points short)."""
    return pymupdf.Font("helv").text_length(text, fontsize=fs)


def test_a_trial_week_fits_the_dates_line_of_the_ucs_form_and_the_mofr(tmp_path):
    """Five days in a row on the UCS form and the MOFR: one range, on one line, at a size that fits; the file
    name still gives the first and last day."""
    s = pat_settings()
    case = make_case({**ROE, "dates": FIVE}, attorneys=[ALEX])
    path = fill(case, ALEX, s, tmp_path, dated=True)
    assert "(9-28-2026 to 10-2-2026)" in path.name and case.get("dates") == FIVE  # the case keeps the list
    (text, fs, rect), (second, _, _) = _dates_fields(path, lambda n: ucs_map.WIDGETS.get(n, n))
    assert text == "9/28/2026–10/2/2026" and second == "" and fs >= 7
    assert _width(text, fs) <= rect.width - 4
    mofr = fill_mofr(case, s, tmp_path)
    with pymupdf.open(mofr) as doc:
        w = next(w for w in doc[0].widgets() if w.field_name == "Text Field9")
        assert w.field_value == "9/28/2026–10/2/2026" and _width(w.field_value, w.text_fontsize) <= w.rect.width - 4


@pytest.mark.parametrize("choice, fmap", [("ucs", ucs_map), ("original", original_map)])
def test_dates_that_do_not_fit_their_line_go_on_under_it(tmp_path, choice, fmap):
    """Five days with gaps: too wide for the line at a readable size, so they are split between dates over the
    line and the one added under it, both at the same size, nothing cut off."""
    s = pat_settings()
    s.form_choice = choice
    case = make_case({**ROE, "dates": GAPS}, attorneys=[ALEX])
    path = fill(case, ALEX, s, tmp_path)
    (text, fs, rect), (second, fs2, rect2) = _dates_fields(
        path, lambda n: fmap.WIDGETS.get(n.removeprefix("Text-"), n))
    assert text and second and f"{text}, {second}" == GAPS
    assert fs == fs2 >= 7
    assert _width(text, fs) <= rect.width - 4 and _width(second, fs2) <= rect2.width - 4
    assert rect2.y0 >= rect.y1 - 1 and rect2.x1 <= rect.x1 + 1  # the second line sits under the first


def test_the_agreement_date_filled_in_is_a_default_not_typed(tmp_path):
    """The date of agreement set to today by Settings is marked as a default: the job's own field, which the
    agreement shares, must not look typed by the user, or every re-merge would keep it."""
    s = pat_settings()
    s.agreement_today = True
    case = make_case(ROE, attorneys=[ALEX])
    fill(case, ALEX, s, tmp_path)
    assert case.get("agreement_date") and case.fields["agreement_date"].source == SRC_DEFAULT
    with pymupdf.open(form_path("ucs")) as blank:
        assert blank.page_count == 1 + ucs_map.INSTRUCTION_PAGES
    s.include_instructions = False
    with pymupdf.open(fill(case, ALEX, s, tmp_path / "x")) as doc:
        assert doc.page_count == 1


# ================================================================== the review of the sweep's own fixes

WEEK = "10/12/2026, 10/13/2026, 10/14/2026, 10/15/2026, 10/16/2026"  # one week of trial in October
WEEK_AND_A_DAY = FIVE + ", 10/5/2026"                               # "9/28/2026–10/2/2026, 10/5/2026"


def _drawn(page, rect) -> str:
    """The text drawn inside a field's rectangle, the form's own underline left out."""
    return " ".join(page.get_text(clip=rect).replace("_", " ").split())


def _key_of(choice):
    """Our key ('dates', 'dates_2') for a field name of this form."""
    fmap = {"ucs": ucs_map, "original": original_map}.get(choice)
    return (lambda n: fmap.WIDGETS.get(n.removeprefix("Text-"), n)) if fmap else (lambda n: n)


@pytest.mark.parametrize("choice", ["ucs", "original", "clean"])
@pytest.mark.parametrize("dates, last", [(WEEK, "10/16/2026"), (WEEK_AND_A_DAY, "10/5/2026")])
def test_a_range_of_dates_is_drawn_to_its_last_day(tmp_path, choice, dates, last):
    """The en dash of a range was measured as if it were narrow, so "10/12/2026–10/16/2026" got too large a size
    and was drawn as "10/12/2026–10/16/202". Every line drawn is now its whole value, the last date at its end."""
    from minute_filler.fill import date_ranges
    s = pat_settings()
    s.form_choice = choice
    path = fill(make_case({**ROE, "dates": dates}, attorneys=[ALEX]), ALEX, s, tmp_path)
    key_of = _key_of(choice)
    with pymupdf.open(path) as doc:
        page = doc[0]
        lines = sorted((w for w in page.widgets() if key_of(w.field_name) in ("dates", "dates_2") and w.field_value),
                       key=lambda w: w.rect.y0)
        for w in lines:
            assert _width(w.field_value, w.text_fontsize) <= w.rect.width - 4
            assert _drawn(page, w.rect) == w.field_value
        drawn = " ".join(_drawn(page, w.rect) for w in lines)
    assert drawn.endswith(last)
    assert re.sub(r"[,\s]+", " ", drawn) == re.sub(r"[,\s]+", " ", date_ranges(dates))


@pytest.mark.parametrize("dates, shown", [
    (WEEK, "10/12/2026–10/16/2026"),
    (WEEK + ", 10/19/2026", "10/12/2026–10/16/2026, 10/19/2026"),
    (WEEK + ", 10/19/2026, 10/20/2026, 10/21/2026, 10/22/2026, 10/23/2026",
     "10/12/2026–10/16/2026, 10/19/2026–10/23/2026"),
])
def test_trial_weeks_in_october_are_drawn_whole_on_the_mofr(tmp_path, dates, shown):
    """The MOFR's dates line shrinks the ranges until they fit: the last day is not cut off ("...10/23/202")."""
    mofr = fill_mofr(make_case({**ROE, "dates": dates}, attorneys=[ALEX]), pat_settings(), tmp_path)
    with pymupdf.open(mofr) as doc:
        w = next(w for w in doc[0].widgets() if w.field_name == "Text Field9")
        assert w.field_value == shown
        assert _width(w.field_value, w.text_fontsize) <= w.rect.width - 4
        assert _drawn(doc[0], w.rect) == w.field_value


@pytest.mark.parametrize("choice", ["ucs", "original"])
@pytest.mark.parametrize("dates, first, second", [
    ("September 28, 2026 and October 5, 2026", "September 28, 2026", "October 5, 2026"),
    ("September 28, 2026, September 30, 2026 and October 5, 2026",
     "September 28, 2026", "September 30, 2026 and October 5, 2026"),
    ("6/3/2026 (a.m. session), 6/4/2026 (p.m. session), 6/5/2026",
     "6/3/2026 (a.m. session)", "6/4/2026 (p.m. session), 6/5/2026"),
])
def test_dates_written_out_go_on_to_the_second_line_between_dates(tmp_path, choice, dates, first, second):
    """Dates typed out are split only between two dates, never at the comma inside one ("September 28" |
    "2026 and October 5, 2026"), and each line keeps the text as typed."""
    s = pat_settings()
    s.form_choice = choice
    path = fill(make_case({**ROE, "dates": dates}, attorneys=[ALEX]), ALEX, s, tmp_path)
    (text, fs, _), (more, fs2, _) = _dates_fields(path, _key_of(choice))
    assert (text, more) == (first, second) and fs == fs2


# ------------------------------------------------------------------ a file left behind, a copy's date

@pytest.fixture
def logfile(tmp_path, monkeypatch):
    """A fresh log in a temporary folder; the path of its file."""
    monkeypatch.setenv("APPDATA", str(tmp_path / "logged"))
    log.shutdown()
    log.setup()
    yield log.log_path()
    log.shutdown()


def _log_text(path) -> str:
    """Everything written to the log so far."""
    for h in log.log.handlers:
        h.flush()
    return path.read_text(encoding="utf-8")


def test_a_half_written_invoice_that_cannot_be_deleted_keeps_its_number(s, tmp_path, monkeypatch, logfile):
    """The save fails part way and the half-written PDF can't be deleted (an antivirus scan holds it): it stays
    on disk, so its number stays taken (the next invoice is not "Invoice 2026-0001 (2).pdf"), and the log names
    neither the case nor the attorney."""
    case = make_case({**ROE, "est_pages": "30"}, attorneys=[ALEX])
    ledger = Ledger(tmp_path / "r.db")
    real_unlink = Path.unlink

    def half(self, path, *args, **kwargs):
        Path(path).write_bytes(b"%PDF-1.7 half of it")
        raise OSError(28, "No space left on device")

    def held(self, missing_ok=False):
        if self.suffix == ".pdf":
            raise PermissionError(13, "The process cannot access the file", str(self))
        return real_unlink(self, missing_ok=missing_ok)
    monkeypatch.setattr(pymupdf.Document, "save", half)
    monkeypatch.setattr(Path, "unlink", held)
    with pytest.raises(OSError, match="No space left"):
        make_invoice(case, ALEX, s, tmp_path / "out", InvoiceOpts(30), ledger)
    year = date.today().year
    left, = (tmp_path / "out").glob("*.pdf")
    assert f"{year}-0001" in left.name
    assert ledger.next_invoice_no()[0] == f"{year}-0002"  # (before: 0001 again, beside the file left behind)
    logged = _log_text(logfile)
    assert "stays taken" in logged
    for private in ("Roe", "Holding", "712345", "Counsel"):
        assert private not in logged


def test_the_detailed_copy_is_dated_as_its_invoice_across_midnight(s, tmp_path, monkeypatch):
    """The invoice made at 11:59 p.m. on December 31 and its detailed copy a moment later: both are dated as the
    invoice and its record, not the copy a day (and a year) later."""
    class Midnight(date):
        calls = 0

        @classmethod
        def today(cls):
            cls.calls += 1
            return date(2026, 12, 31) if cls.calls == 1 else date(2027, 1, 1)
    for module in (invoice, records, deliver):
        monkeypatch.setattr(module, "date", Midnight, raising=False)
    s.invoice_detailed_copy = True
    case = make_case({**ROE, "dates": "12/30/2026", "est_pages": "30"}, attorneys=[ALEX])
    ledger = Ledger(tmp_path / "r.db")
    made = deliver.generate(case, s, tmp_path / "out", ["invoice"], InvoiceOpts(30), ledger)
    dated = {}
    for p in made:
        with pymupdf.open(p) as doc:
            fields = {w.field_name: w.field_value for w in doc[0].widgets()}
        dated["copy" if "(detailed)" in p.name else "invoice"] = (fields["invoice_no"], fields["date"])
    inv, = ledger.invoices()
    assert dated["invoice"] == dated["copy"] == (inv.invoice_no, "12/31/2026") and inv.created == "2026-12-31"


# ------------------------------------------------------------------ run sheets and the log

def test_a_workbook_of_the_case_that_is_no_run_sheet_does_not_hold_up_its_run_sheet(tmp_path, s):
    """Next to the transcript, "712345-2021 Roe v Poe exhibit list.xlsx" is password-protected (no zip) and a
    damaged workbook has the index number too: openpyxl can't open them, now or later. They are not run sheets,
    so the case's own run sheet is found and added to, instead of the error stopping it for good."""
    case = make_case(ROE)
    first = RunSheetOpts([Row(date(2026, 6, 3), "Pat", 20, 1)])
    add_takes(case, first, s, [])
    docs = tmp_path / "Roe v Poe"
    docs.mkdir()
    (docs / "712345-2021 Roe v Poe exhibit list.xlsx").write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 600)
    with zipfile.ZipFile(docs / "712345-2021 Roe v Poe damaged.xlsx", "w") as z:
        z.writestr("notes.txt", "not a workbook")
    assert [f.path for f in find_sheets(case, [tmp_path / "sheets", docs])] == [first.path]
    again = RunSheetOpts([Row(date(2026, 6, 4), "Pat", 5, 21)])
    assert add_takes(case, again, s, [docs]) == first.path and not again.created
    assert [p.name for p in (tmp_path / "sheets").glob("*.xlsx")] == [first.path.name]  # no "(2)"


def test_a_workbook_of_the_case_held_by_another_program_still_stops_the_run_sheet(tmp_path, s, monkeypatch):
    """Opening the case's run sheet fails with an OSError (a sync holds it): that can pass, so it is still
    SheetUnreadable, naming the file, and no second run sheet is started."""
    case = make_case(ROE)
    first = RunSheetOpts([Row(date(2026, 6, 3), "Pat", 20, 1)])
    add_takes(case, first, s, [])
    real = openpyxl.load_workbook

    def synced(path, *args, **kwargs):
        if Path(path).name == first.path.name:
            raise OSError(22, "The cloud file provider is not running")
        return real(path, *args, **kwargs)
    monkeypatch.setattr(openpyxl, "load_workbook", synced)
    with pytest.raises(SheetUnreadable, match=re.escape(first.path.name)):
        find_sheets(case, [tmp_path / "sheets"])
    with pytest.raises(SheetUnreadable):
        add_takes(case, RunSheetOpts([Row(date(2026, 6, 4), "Pat", 5, 21)]), s, [])
    assert [p.name for p in (tmp_path / "sheets").glob("*.xlsx")] == [first.path.name]


def test_the_log_names_no_case_of_a_file_it_could_not_read(tmp_path, logfile):
    """A run sheet that can't be read, by its name, and an error whose message names the file: the log gets
    neither the case nor its index number (the log goes out with problem reports)."""
    name = "June 2026 712345-2021 Jane Roe v. Sam Poe - Run Sheet.xlsx"
    bad = tmp_path / name
    bad.write_bytes(b"not a zip")
    assert read_info(bad) is None
    try:
        raise ValueError("File is not a zip file")
    except ValueError as e:
        log.error(f"could not read {name}", e)
    logged = _log_text(logfile)
    assert "not a run sheet" in logged and "<file.xlsx>" in logged  # (log.error scrubs its message too)
    for private in ("Roe", "Poe", "712345"):
        assert private not in logged


# ------------------------------------------------------------------ the word index

def _index(entry) -> str:
    """The first page of a word index: its numbers first, in the concordance's order, then words."""
    numbers = ["1", "10", "100", "11", "12", "13", "14", "15", "16", "17", "18", "19", "2", "20", "2026"]
    return "\n".join([entry(n) for n in numbers] + ["able [3] 12:4 15:6 20:1", "about [12] 3:4 5:6 7:8"])


def test_a_word_index_whose_numbers_run_10_to_19_is_not_a_page_of_testimony(tmp_path, s):
    """In the index, "10" ... "19" follow each other at the left edge like line numbers. They don't start at 1, as
    a page's lines do, so the page is still the index: not counted, not billed. Pages of testimony still count:
    lines 1 to 25, a page number above them, a short last page."""
    for entry in (lambda n: f"{n} [2] 45:3 67:12", lambda n: f"{n} 45:3, 67:12", lambda n: f"{n}\n(2) 45:3;67:12"):
        assert not has_line_numbers(_index(entry)) and is_index_page(_index(entry))
    lines = [f"{i:2}    Q.  And at 10:{i:02} what did the chart say?" for i in range(1, 26)]
    assert has_line_numbers("\n".join(lines))
    assert has_line_numbers("312\n" + "\n".join(lines))
    assert has_line_numbers("\n".join(lines[:9]) + "\n(Whereupon, the trial was adjourned.)")
    assert not has_line_numbers("\n".join(lines[9:20]))  # lines 10 to 20 alone: no page starts there
    pages = DAY_1[:-1] + [(None, "", "Roe v. Poe", _index(lambda n: f"{n} [2] 45:3 67:12"))]
    path = transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf", pages)
    with pymupdf.open(path) as doc:
        assert len(scan_pdf(doc)) == 10
    assert job_of(path, s).transcript_pages() == 10


def test_a_date_written_out_names_the_file_with_its_year(s):
    """The Dates field typed 'September 28, 2026 and October 5, 2026': {date} in the file name is '9-28-2026', not
    'September 28' (cut at the comma inside the date)."""
    from minute_filler.fill import output_name
    case = make_case({"case_name": "Jane Roe v. Sam Poe", "index_no": "712345/2021",
                      "dates": "September 28, 2026 and October 5, 2026"})
    s.filename_pattern = "{date} {case}"
    assert output_name(case, None, s) == "9-28-2026 Jane Roe v. Sam Poe.pdf"
    case.set("dates", "9/28/2026, 10/5/2026")
    assert output_name(case, None, s) == "9-28-2026 Jane Roe v. Sam Poe.pdf"
