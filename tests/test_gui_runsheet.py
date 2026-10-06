"""The run sheet from the window (off screen): Generate and Generate all ask where the takes go when the case
has one (once per trial), one run sheet counted per case, a run sheet open in Excel, the Settings → Run sheet tab
(initials and reporter lines written different ways) and its folder, and work that ends while a question is on
screen."""
import os
from copy import deepcopy

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
QtWidgets = pytest.importorskip("PySide6.QtWidgets")

from openpyxl import load_workbook  # noqa: E402

from helpers import pat_settings  # noqa: E402
from test_gui_batch import wait  # noqa: E402
from test_runsheet import transcript  # noqa: E402


@pytest.fixture
def window(tmp_path, monkeypatch, make_window):
    monkeypatch.setattr(QtWidgets.QMessageBox, "exec", lambda self: 0)  # no dialogs waiting for a click
    s = pat_settings()
    s.use_ai = s.open_after = False
    s.output_dir = str(tmp_path / "out")
    s.runsheet_dir = str(tmp_path / "sheets")
    s.outputs = ["runsheet"]
    return make_window(s)


def takes_on(path):
    """The Reporter column of a run sheet, one name per take."""
    ws =load_workbook(path)["Run Sheet"]
    return [ws.cell(r, 3).value for r in range(5, ws.max_row + 1)]


def test_generate_starts_a_run_sheet_then_asks(window, tmp_path, monkeypatch):
    from minute_filler.gui.dialogs import RunSheetDialog
    window.add_files([str(transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf"))])
    wait(window, lambda: len(window.cur.docs) == 1)
    assert window.output_boxes["runsheet"].isChecked() and "takes" in window.output_boxes["runsheet"].toolTip()
    asked = []
    monkeypatch.setattr(RunSheetDialog, "exec", lambda self: asked.append(self) or 0)  # Cancel
    window.fill()  # no run sheet yet: nothing to ask, and the attorneys are not asked about for a run sheet
    sheet, = (tmp_path / "sheets").glob("*.xlsx")
    assert takes_on(sheet) == ["Pat", "Dana", "Pat", "Dana"] and not asked and window.cur.saved == [sheet]

    window.fill()  # now there is one: asked, and Cancel makes nothing
    assert len(asked) == 1 and list((tmp_path / "sheets").glob("*.xlsx")) == [sheet]
    monkeypatch.setattr(RunSheetDialog, "exec", lambda self: RunSheetDialog.Accepted)  # add to it
    window.fill()
    assert len(takes_on(sheet)) == 4 and window.cur.saved == []  # the same takes are not added twice

    window.s.runsheet_existing = "new"
    window.fill()
    assert len(list((tmp_path / "sheets").glob("*.xlsx"))) == 2


def test_run_sheet_settings(qt, tmp_path, monkeypatch):
    """Initials and reporter names are tidied on Save, a line that can't be read is named, and they show
    again the same way."""
    from minute_filler.gui.dialogs import SettingsDialog
    warned = []
    monkeypatch.setattr(QtWidgets.QMessageBox, "warning", lambda *a: warned.append(a[2]))
    s = pat_settings()
    dlg = SettingsDialog(s)
    dlg.p_initials.setText("P.R.")
    dlg.r_names.setPlainText("DS = Dana\nkl: Kim\nnonsense\n")
    dlg.r_existing.setCurrentIndex(dlg.r_existing.findData("add"))
    dlg.accept()
    assert s.profile.initials == "pr" and s.reporter_names() == {"ds": "Dana", "kl": "Kim"}
    assert len(warned) == 1 and "nonsense" in warned[0]  # the rest is saved; the line left out is named
    assert s.runsheet_existing == "add"
    again = SettingsDialog(s)
    assert again.r_names.toPlainText() == "ds = Dana\nkl = Kim"


def test_generate_all_asks_once_per_trial(window, tmp_path, monkeypatch):
    from minute_filler.gui.dialogs import RunSheetDialog
    from test_runsheet import TALK, TITLE_1
    first = transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf")
    window.add_files([str(first)])
    wait(window, lambda: len(window.cur.docs) == 1)
    window.fill()  # the run sheet of the trial so far
    second = transcript(tmp_path / "Transcript 6-4-2026 Roe v Poe.pdf", [(320, "ds", "", TITLE_1), (321, "pr", "", TALK)],
                        day="June 4, 2026")
    third = transcript(tmp_path / "Transcript 6-5-2026 Roe v Poe.pdf", [(322, "pr", "", TITLE_1), (323, "ds", "", TALK)],
                       day="June 5, 2026")
    window.new_job()
    window.add_files([str(second), str(third)])
    wait(window, lambda: len(window.jobs) == 2 and all(j.docs for j in window.jobs))
    asked = []
    monkeypatch.setattr(RunSheetDialog, "exec", lambda self: asked.append(self) or RunSheetDialog.Accepted)
    window.fill_all_jobs()
    wait(window, lambda: not window.filling)
    sheet, = (tmp_path / "sheets").glob("*.xlsx")
    assert len(asked) == 1 and takes_on(sheet) == ["Pat", "Dana", "Pat", "Dana", "Dana", "Pat", "Pat", "Dana"]


# ------------------------------------------------------------------ found in the run sheet sweep

def two_days(tmp_path, window):
    """The case's run sheet so far (June 3), then a batch of two more days (June 4 and 5): two jobs."""
    from test_runsheet import TALK, TITLE_1
    window.add_files([str(transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf"))])
    wait(window, lambda: len(window.cur.docs) == 1)
    window.fill()
    second = transcript(tmp_path / "Transcript 6-4-2026 Roe v Poe.pdf", [(320, "ds", "", TITLE_1), (321, "pr", "", TALK)],
                        day="June 4, 2026")
    third = transcript(tmp_path / "Transcript 6-5-2026 Roe v Poe.pdf", [(322, "pr", "", TITLE_1), (323, "ds", "", TALK)],
                       day="June 5, 2026")
    window.new_job()
    window.add_files([str(second), str(third)])
    wait(window, lambda: len(window.jobs) == 2 and all(j.docs for j in window.jobs))
    sheet, = (tmp_path / "sheets").glob("*.xlsx")
    return sheet


def test_files_follow_an_answer_that_came_in_during_the_question(window, tmp_path, monkeypatch):
    """The AI's answer (or a read) that ends while the run sheet question is on screen changes the job: the
    files are made from the job as it is then, as the window shows it."""
    from minute_filler.gui.dialogs import RunSheetDialog
    window.add_files([str(transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf"))])
    wait(window, lambda: len(window.cur.docs) == 1)
    window.fill()  # the case has a run sheet now: the next Generate asks

    def answer_comes_in(dlg):
        case = deepcopy(window.cur.case)
        case.set("case_name", "Jane Roe v. Terry Hill")
        window.cur.case = case
        return RunSheetDialog.Accepted
    monkeypatch.setattr(RunSheetDialog, "exec", answer_comes_in)
    monkeypatch.setattr(RunSheetDialog, "choice", lambda self: "")  # start a new run sheet
    window.fill()
    names = [p.name for p in (tmp_path / "sheets").glob("*.xlsx")]
    assert len(names) == 2 and any("Terry Hill" in n for n in names)


def test_jobs_joined_while_asking_about_the_run_sheet(window, tmp_path, monkeypatch):
    """A job that is gone when the question closes is left out of Generate all; the rest is made."""
    from minute_filler.gui.dialogs import RunSheetDialog
    sheet = two_days(tmp_path, window)

    def joined(dlg):  # a read that ended meanwhile put the second job's documents into the first
        window.jobs.remove(window.jobs[1])
        return RunSheetDialog.Accepted
    monkeypatch.setattr(RunSheetDialog, "exec", joined)
    window.fill_all_jobs()
    wait(window, lambda: not window.filling)
    assert takes_on(sheet) == ["Pat", "Dana", "Pat", "Dana", "Dana", "Pat"]


def test_ctrl_enter_while_typing_in_the_attorney_table(window, monkeypatch):
    """Generate takes the attorney cell still being typed in as typed."""
    told = []
    monkeypatch.setattr(QtWidgets.QMessageBox, "information", lambda *a: told.append(a[2]))
    window.show()
    window.activateWindow()
    window._add_att_row()  # the cursor is in the new row's Name cell
    QtWidgets.QApplication.processEvents()
    editor = QtWidgets.QApplication.focusWidget()
    assert isinstance(editor, QtWidgets.QLineEdit)
    editor.setText("Alex B. Counsel")
    window.s.outputs = []
    window.fill()  # Ctrl+Enter: nothing ticked to make, but what was typed is kept
    assert [a.name for a in window.cur.case.attorneys] == ["Alex B. Counsel"]
    assert "run sheet" in told[0]


def test_generate_all_counts_one_run_sheet_per_case(window, tmp_path):
    two_days(tmp_path, window)
    window._refresh_job_labels()
    assert window.fill_all_btn.text() == "Generate all  (1 file)"


def test_batch_message_says_why_the_run_sheet_was_not_saved(window, tmp_path, monkeypatch):
    """A run sheet open in Excel: the batch's message says so plainly, without the Python error name."""
    import minute_filler.runsheet as rs
    from test_runsheet import TALK, TITLE_1
    window.add_files([str(transcript(tmp_path / "Transcript 6-3-2026 Roe v Poe.pdf")),
                      str(transcript(tmp_path / "Transcript 6-4-2026 Roe v Poe.pdf", [(320, "ds", "", TITLE_1),
                                     (321, "pr", "", TALK)], day="June 4, 2026"))])
    wait(window, lambda: len(window.jobs) == 2 and all(j.docs for j in window.jobs))

    def locked(wb, path):
        raise PermissionError(f"{path.name} is open in another program - close it in Excel and try again")
    monkeypatch.setattr(rs, "_save", locked)
    shown = []
    monkeypatch.setattr(window, "_saved_box", lambda title, text, *a, **kw: shown.append(text))
    window.fill_all_jobs()
    wait(window, lambda: not window.filling)
    assert "Excel" in shown[0].split("\n\n")[1]  # the heading names the run sheet too
    assert "Run Sheet.xlsx is open in another program - close it in Excel" in shown[0]
    assert "PermissionError" not in shown[0]


def test_reading_gets_its_own_copy_of_the_settings(window, monkeypatch):
    """The background read works on a copy of the settings: Settings saved meanwhile can't change them under it."""
    import minute_filler.gui.main_window as mw
    seen = []
    real = mw.read_loaders
    monkeypatch.setattr(mw, "read_loaders", lambda loaders, names, s, *a: seen.append(s) or real(loaders, names, s, *a))
    window.add_text("Please send the minutes in Smith v Jones, Index No. 712222/2024, 5/22/2026.")
    wait(window, lambda: bool(window.cur.docs))
    assert seen and seen[0] is not window.s and seen[0].profile.name == "Pat Reporter"


def test_reporter_lines_written_other_ways(qt, monkeypatch):
    """"D. S. = Dana", "kl Kim", "mt - Maria" and a tab are all read; "Dana Smith - ds" (name first) is named."""
    from minute_filler.gui.dialogs import SettingsDialog
    warned = []
    monkeypatch.setattr(QtWidgets.QMessageBox, "warning", lambda *a: warned.append(a[2]))
    s = pat_settings()
    dlg = SettingsDialog(s)
    dlg.r_names.setPlainText("D. S. = Dana\nkl Kim\nmt - Maria\npr\tPat\nDana Smith - ds\n")
    dlg.accept()
    assert s.reporter_names() == {"ds": "Dana", "kl": "Kim", "mt": "Maria", "pr": "Pat"}
    assert "Dana Smith - ds" in warned[0]


def test_open_the_run_sheets_folder_before_there_is_one(qt, tmp_path, monkeypatch):
    import minute_filler.gui.dialogs as dialogs
    opened = []
    monkeypatch.setattr(dialogs, "open_path", opened.append)
    dlg = dialogs.SettingsDialog(pat_settings())
    edit = dlg.o_dirs["runsheet"]  # Options -> Folders
    edit.setText(str(tmp_path / "Run Sheets"))
    dlg._open_folder(edit, "runsheet")
    assert (tmp_path / "Run Sheets").is_dir() and opened == [tmp_path / "Run Sheets"]
